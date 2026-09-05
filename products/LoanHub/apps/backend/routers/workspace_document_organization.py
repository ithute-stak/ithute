from __future__ import annotations

import html
from io import BytesIO
from pathlib import Path
from uuid import UUID

from docx import Document as WordDocument
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile, status
from pydantic import BaseModel, Field
from pypdf import PdfReader
from sqlalchemy import func
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from core.access_control import TenantContext, get_user_context, is_platform_role
from database.models.workspace_document import WorkspaceDocument, WorkspaceDocumentRevision
from database.models.workspace_document_folder import WorkspaceDocumentFolder, WorkspaceDocumentFolderItem
from database.session import get_db
from services.workspace_document_service import (
    document_reference,
    plain_text_from_html,
    sanitize_document_html,
)


router = APIRouter(prefix="/workspace-documents", tags=["Document Studio"])
MAX_IMPORT_BYTES = 20 * 1024 * 1024
SUPPORTED_IMPORTS = {".docx", ".pdf"}


class FolderCreate(BaseModel):
    name: str = Field(min_length=1, max_length=120)
    parent_id: UUID | None = None


class FolderRename(BaseModel):
    name: str = Field(min_length=1, max_length=120)


class FolderRead(BaseModel):
    id: UUID
    name: str
    parent_id: UUID | None = None
    document_count: int = 0
    child_count: int = 0


class FolderMove(BaseModel):
    folder_id: UUID | None = None


class ImportedDocumentRead(BaseModel):
    id: UUID
    reference: str
    title: str
    folder_id: UUID | None = None
    source_type: str
    warning: str | None = None


def _normalise_folder_name(value: str) -> tuple[str, str]:
    name = " ".join(value.split()).strip()
    if not name:
        raise HTTPException(status_code=422, detail="Folder name cannot be blank")
    return name, name.casefold()


def _folder_or_404(db: Session, context: TenantContext, folder_id: UUID) -> WorkspaceDocumentFolder:
    folder = db.query(WorkspaceDocumentFolder).filter(
        WorkspaceDocumentFolder.id == folder_id,
        WorkspaceDocumentFolder.owner_user_id == context.user.id,
        WorkspaceDocumentFolder.is_deleted.is_(False),
    ).first()
    if not folder:
        raise HTTPException(status_code=404, detail="Folder not found")
    return folder


def _owned_document_or_404(db: Session, context: TenantContext, document_id: UUID) -> WorkspaceDocument:
    document = db.query(WorkspaceDocument).filter(
        WorkspaceDocument.id == document_id,
        WorkspaceDocument.owner_user_id == context.user.id,
        WorkspaceDocument.is_deleted.is_(False),
    ).first()
    if not document:
        raise HTTPException(status_code=404, detail="Document not found")
    return document


def _validate_parent(db: Session, context: TenantContext, parent_id: UUID | None, folder_id: UUID | None = None) -> WorkspaceDocumentFolder | None:
    if parent_id is None:
        return None
    parent = _folder_or_404(db, context, parent_id)
    if folder_id and parent.id == folder_id:
        raise HTTPException(status_code=400, detail="A folder cannot be its own parent")
    current = parent
    seen: set[UUID] = set()
    while current.parent_id:
        if current.id in seen:
            break
        seen.add(current.id)
        if folder_id and current.parent_id == folder_id:
            raise HTTPException(status_code=400, detail="Cannot move a folder inside one of its descendants")
        current = _folder_or_404(db, context, current.parent_id)
    return parent


def _folder_read(db: Session, folder: WorkspaceDocumentFolder) -> FolderRead:
    document_count = db.query(func.count(WorkspaceDocumentFolderItem.id)).filter(
        WorkspaceDocumentFolderItem.folder_id == folder.id
    ).scalar() or 0
    child_count = db.query(func.count(WorkspaceDocumentFolder.id)).filter(
        WorkspaceDocumentFolder.parent_id == folder.id,
        WorkspaceDocumentFolder.is_deleted.is_(False),
    ).scalar() or 0
    return FolderRead(
        id=folder.id,
        name=folder.name,
        parent_id=folder.parent_id,
        document_count=int(document_count),
        child_count=int(child_count),
    )


@router.get("/folders", response_model=list[FolderRead])
def list_folders(
    db: Session = Depends(get_db),
    context: TenantContext = Depends(get_user_context),
):
    rows = db.query(WorkspaceDocumentFolder).filter(
        WorkspaceDocumentFolder.owner_user_id == context.user.id,
        WorkspaceDocumentFolder.is_deleted.is_(False),
    ).order_by(WorkspaceDocumentFolder.name.asc()).all()
    return [_folder_read(db, row) for row in rows]


@router.post("/folders", response_model=FolderRead, status_code=status.HTTP_201_CREATED)
def create_folder(
    payload: FolderCreate,
    db: Session = Depends(get_db),
    context: TenantContext = Depends(get_user_context),
):
    name, normalized_name = _normalise_folder_name(payload.name)
    _validate_parent(db, context, payload.parent_id)
    row = WorkspaceDocumentFolder(
        owner_user_id=context.user.id,
        company_id=context.company_id,
        parent_id=payload.parent_id,
        name=name,
        normalized_name=normalized_name,
    )
    db.add(row)
    try:
        db.commit()
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(status_code=409, detail="A folder with that name already exists here") from error
    db.refresh(row)
    return _folder_read(db, row)


@router.patch("/folders/{folder_id}", response_model=FolderRead)
def rename_folder(
    folder_id: UUID,
    payload: FolderRename,
    db: Session = Depends(get_db),
    context: TenantContext = Depends(get_user_context),
):
    row = _folder_or_404(db, context, folder_id)
    name, normalized_name = _normalise_folder_name(payload.name)
    row.name = name
    row.normalized_name = normalized_name
    try:
        db.commit()
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(status_code=409, detail="A folder with that name already exists here") from error
    db.refresh(row)
    return _folder_read(db, row)


@router.delete("/folders/{folder_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_folder(
    folder_id: UUID,
    db: Session = Depends(get_db),
    context: TenantContext = Depends(get_user_context),
):
    row = _folder_or_404(db, context, folder_id)
    child_count = db.query(func.count(WorkspaceDocumentFolder.id)).filter(
        WorkspaceDocumentFolder.parent_id == row.id,
        WorkspaceDocumentFolder.is_deleted.is_(False),
    ).scalar() or 0
    if child_count:
        raise HTTPException(status_code=409, detail="Move or delete the subfolders first")
    # Documents are not deleted with a folder. They simply return to My documents.
    db.query(WorkspaceDocumentFolderItem).filter(
        WorkspaceDocumentFolderItem.folder_id == row.id
    ).delete(synchronize_session=False)
    row.is_deleted = True
    db.commit()


@router.patch("/{document_id}/folder", status_code=status.HTTP_204_NO_CONTENT)
def move_document(
    document_id: UUID,
    payload: FolderMove,
    db: Session = Depends(get_db),
    context: TenantContext = Depends(get_user_context),
):
    _owned_document_or_404(db, context, document_id)
    if payload.folder_id:
        _folder_or_404(db, context, payload.folder_id)
    assignment = db.query(WorkspaceDocumentFolderItem).filter(
        WorkspaceDocumentFolderItem.document_id == document_id
    ).first()
    if payload.folder_id is None:
        if assignment:
            db.delete(assignment)
    elif assignment:
        assignment.folder_id = payload.folder_id
    else:
        db.add(WorkspaceDocumentFolderItem(folder_id=payload.folder_id, document_id=document_id))
    db.commit()


@router.get("/folder-assignments", response_model=dict[str, str])
def folder_assignments(
    db: Session = Depends(get_db),
    context: TenantContext = Depends(get_user_context),
):
    rows = db.query(WorkspaceDocumentFolderItem).join(
        WorkspaceDocumentFolder,
        WorkspaceDocumentFolder.id == WorkspaceDocumentFolderItem.folder_id,
    ).join(
        WorkspaceDocument,
        WorkspaceDocument.id == WorkspaceDocumentFolderItem.document_id,
    ).filter(
        WorkspaceDocumentFolder.owner_user_id == context.user.id,
        WorkspaceDocumentFolder.is_deleted.is_(False),
        WorkspaceDocument.owner_user_id == context.user.id,
        WorkspaceDocument.is_deleted.is_(False),
    ).all()
    return {str(row.document_id): str(row.folder_id) for row in rows}


def _docx_html(data: bytes) -> str:
    document = WordDocument(BytesIO(data))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            parts.append("<p></p>")
            continue
        style = (paragraph.style.name or "").lower() if paragraph.style else ""
        tag = "p"
        if style.startswith("heading"):
            try:
                level = max(1, min(4, int(style.split()[-1])))
                tag = f"h{level}"
            except (ValueError, IndexError):
                tag = "h2"
        parts.append(f"<{tag}>{html.escape(text)}</{tag}>")
    for table in document.tables:
        rows: list[str] = []
        for row in table.rows:
            cells = "".join(f"<td>{html.escape(cell.text.strip())}</td>" for cell in row.cells)
            rows.append(f"<tr>{cells}</tr>")
        if rows:
            parts.append(f"<table><tbody>{''.join(rows)}</tbody></table>")
    return "".join(parts) or "<p></p>"


def _pdf_html(data: bytes) -> tuple[str, str | None]:
    try:
        reader = PdfReader(BytesIO(data))
    except Exception as error:
        raise HTTPException(status_code=422, detail="The PDF could not be read") from error
    if reader.is_encrypted:
        raise HTTPException(status_code=422, detail="Password-protected PDFs must be unlocked before import")
    page_text: list[str] = []
    for page in reader.pages:
        text = (page.extract_text() or "").strip()
        if text:
            page_text.append(text)
    if not page_text:
        raise HTTPException(
            status_code=422,
            detail="This PDF appears to contain only scanned images. Run OCR or export it as a text PDF before importing for editing.",
        )
    blocks = []
    for page_index, text in enumerate(page_text):
        if page_index:
            blocks.append("<div data-page-break='true'></div>")
        for paragraph in [item.strip() for item in text.split("\n") if item.strip()]:
            blocks.append(f"<p>{html.escape(paragraph)}</p>")
    return "".join(blocks), "PDF text was imported as editable paragraphs; complex original layout may need adjustment."


def _tiptap_from_html(content_html: str) -> dict:
    plain = plain_text_from_html(content_html)
    paragraphs = [line.strip() for line in plain.splitlines() if line.strip()]
    return {
        "type": "doc",
        "content": [
            {"type": "paragraph", "content": [{"type": "text", "text": line}]}
            for line in paragraphs
        ] or [{"type": "paragraph"}],
    }


@router.post("/import", response_model=ImportedDocumentRead, status_code=status.HTTP_201_CREATED)
async def import_document(
    file: UploadFile = File(...),
    folder_id: UUID | None = Form(default=None),
    visibility: str = Form(default="private"),
    db: Session = Depends(get_db),
    context: TenantContext = Depends(get_user_context),
):
    filename = Path(file.filename or "Imported document")
    extension = filename.suffix.lower()
    if extension not in SUPPORTED_IMPORTS:
        raise HTTPException(status_code=415, detail="Document Studio imports .docx and .pdf files")
    if visibility not in {"private", "company", "platform"}:
        raise HTTPException(status_code=422, detail="Invalid document visibility")
    if visibility == "company" and not context.company_id:
        raise HTTPException(status_code=400, detail="Company visibility requires an active company")
    if visibility == "platform" and not is_platform_role(context.user.role):
        raise HTTPException(status_code=403, detail="Platform visibility requires a platform role")
    if folder_id:
        _folder_or_404(db, context, folder_id)

    data = await file.read(MAX_IMPORT_BYTES + 1)
    if len(data) > MAX_IMPORT_BYTES:
        raise HTTPException(status_code=413, detail="Document imports are limited to 20 MB")
    if not data:
        raise HTTPException(status_code=422, detail="The uploaded document is empty")

    warning = None
    if extension == ".docx":
        try:
            imported_html = _docx_html(data)
        except Exception as error:
            raise HTTPException(status_code=422, detail="The Word document could not be read") from error
    else:
        imported_html, warning = _pdf_html(data)

    safe_html = sanitize_document_html(imported_html)
    title = filename.stem.strip()[:255] or "Imported document"
    document = WorkspaceDocument(
        owner_user_id=context.user.id,
        company_id=context.company_id,
        branch_id=context.branch_id,
        last_edited_by_user_id=context.user.id,
        reference=document_reference(),
        title=title,
        template_key="blank",
        content_json=_tiptap_from_html(safe_html),
        content_html=safe_html,
        plain_text=plain_text_from_html(safe_html),
        visibility=visibility,
        status="draft",
        version=1,
        include_brand_header=False,
        include_footer=True,
        is_confidential=False,
    )
    db.add(document)
    db.flush()
    db.add(WorkspaceDocumentRevision(
        document_id=document.id,
        version=1,
        title=document.title,
        content_json=document.content_json,
        content_html=document.content_html,
        plain_text=document.plain_text,
        created_by_user_id=context.user.id,
    ))
    if folder_id:
        db.add(WorkspaceDocumentFolderItem(folder_id=folder_id, document_id=document.id))
    db.commit()
    db.refresh(document)
    return ImportedDocumentRead(
        id=document.id,
        reference=document.reference,
        title=document.title,
        folder_id=folder_id,
        source_type=extension.lstrip("."),
        warning=warning,
    )
