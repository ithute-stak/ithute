from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path
from typing import Any, Literal
from urllib.parse import quote
from uuid import UUID

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile, status
from fastapi.responses import StreamingResponse
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session, joinedload

from core.access_control import COMPANY_MANAGEMENT_ROLES, TenantContext, get_user_context
from database.models.borrower import Borrower
from database.models.client_loan_company import ClientCompanyLoan
from database.models.payment import PaymentTransaction
from database.models.user import User
from database.models.workspace_document import WorkspaceDocument
from database.schemas.file_management import ManagedFileRead
from database.schemas.workspace_document import WorkspaceDocumentRead
from database.session import get_db
from routers.workspace_documents import (
    _can_edit,
    _commit_created_document,
    _read,
    _revision,
    _validate_visibility,
)
from routers.workspace_office_governance import _governed_document_or_404
from services.file_service import store_bytes
from services.workspace_document_service import document_reference
from services.workspace_spreadsheet_service import (
    SpreadsheetValidationError,
    blank_workbook,
    export_csv,
    export_xlsx,
    import_workbook,
    normalize_workbook,
    workbook_plain_text,
)


router = APIRouter(prefix="/workspace-spreadsheets", tags=["Spreadsheet Studio"])

SpreadsheetTemplate = Literal["blank", "loan_portfolio", "cashbook", "collections", "budget"]
SpreadsheetVisibility = Literal["private", "company", "platform"]
SpreadsheetExportFormat = Literal["xlsx", "csv"]
DATABASE_TEMPLATES = {"loan_portfolio", "cashbook", "collections"}


class WorkspaceSpreadsheetCreate(BaseModel):
    title: str = Field(default="Untitled workbook", min_length=1, max_length=255)
    template: SpreadsheetTemplate = "blank"
    visibility: SpreadsheetVisibility = "private"
    is_confidential: bool = False


class WorkspaceSpreadsheetUpdate(BaseModel):
    workbook: dict[str, Any]
    title: str | None = Field(default=None, min_length=1, max_length=255)
    expected_version: int | None = Field(default=None, ge=1)
    create_revision: bool = True


class WorkspaceSpreadsheetPublish(BaseModel):
    format: SpreadsheetExportFormat = "xlsx"
    visibility: SpreadsheetVisibility | None = None
    file_name: str | None = Field(default=None, max_length=255)
    is_confidential: bool | None = None


def _enum_value(value: Any) -> str:
    return str(getattr(value, "value", value) or "")


def _date_value(value: Any) -> str | None:
    if value is None:
        return None
    return value.isoformat() if hasattr(value, "isoformat") else str(value)


def _borrower_name(loan: ClientCompanyLoan) -> str:
    borrower = loan.borrower
    person = borrower.user.person if borrower and borrower.user else None
    return person.full_name if person else "Borrower"


def _header_cells(headers: list[str], fill: str) -> dict[str, dict[str, Any]]:
    return {
        f"{get_column_letter(index)}1": {
            "value": header,
            "style": {
                "bold": True,
                "fill_color": fill,
                "horizontal": "center",
                "wrap": True,
            },
        }
        for index, header in enumerate(headers, start=1)
    }


def _branch_scope(context: TenantContext) -> UUID | None:
    if not context.staff or context.role in COMPANY_MANAGEMENT_ROLES:
        return None
    return context.branch_id


def _company_data_workbook(
    db: Session,
    context: TenantContext,
    template: str,
) -> dict[str, Any]:
    """Build a workbook from persisted tenant data only.

    The templates never invent demonstration rows. Company owners/admins can
    work across their company, while branch-scoped staff receive only records
    belonging to their active branch.
    """
    workbook = blank_workbook(template=template)
    sheet = workbook["sheets"][0]
    if not context.company_id or template not in DATABASE_TEMPLATES:
        return workbook

    branch_id = _branch_scope(context)

    if template in {"loan_portfolio", "collections"}:
        loan_query = (
            db.query(ClientCompanyLoan)
            .options(
                joinedload(ClientCompanyLoan.borrower)
                .joinedload(Borrower.user)
                .joinedload(User.person),
                joinedload(ClientCompanyLoan.branch),
            )
            .filter(ClientCompanyLoan.company_id == context.company_id)
        )
        if branch_id is not None:
            loan_query = loan_query.filter(ClientCompanyLoan.branch_id == branch_id)
        loans = (
            loan_query.order_by(ClientCompanyLoan.created_at.desc())
            .limit(2000)
            .all()
        )

        if template == "collections":
            headers = [
                "Borrower",
                "Loan",
                "Balance due",
                "Days overdue",
                "Next due",
                "Status",
                "Risk",
                "Branch",
                "Last contact",
                "Promise date",
                "Outcome",
                "Owner",
            ]
            cells = _header_cells(headers, "#FFF4E5")
            overdue_loans = [
                loan
                for loan in loans
                if loan.is_overdue or _enum_value(loan.status) == "defaulted"
            ]
            for row, loan in enumerate(overdue_loans, start=2):
                due_date = loan.first_payment_due
                days_overdue = (
                    max(0, (date.today() - due_date).days)
                    if due_date and due_date < date.today()
                    else 0
                )
                values = [
                    _borrower_name(loan),
                    loan.loan_reference,
                    float(loan.balance or 0),
                    days_overdue,
                    _date_value(due_date),
                    _enum_value(loan.status),
                    _enum_value(loan.risk_level),
                    loan.branch.name if loan.branch else "",
                ]
                for column, value in enumerate(values, start=1):
                    cells[f"{get_column_letter(column)}{row}"] = {"value": value}
            sheet.update(
                {
                    "cells": cells,
                    "row_count": max(40, len(overdue_loans) + 5),
                    "column_count": len(headers),
                    "freeze_panes": "A2",
                    "column_widths": {
                        "A": 24,
                        "B": 20,
                        "C": 16,
                        "D": 14,
                        "E": 15,
                        "F": 14,
                        "G": 12,
                        "H": 18,
                        "I": 18,
                        "J": 16,
                        "K": 22,
                        "L": 20,
                    },
                }
            )
        else:
            headers = [
                "Loan reference",
                "Borrower",
                "Principal",
                "Balance",
                "Installment",
                "Amount paid",
                "Next due",
                "Maturity",
                "Status",
                "Risk",
                "Overdue",
                "Branch",
            ]
            cells = _header_cells(headers, "#EAF2FF")
            for row, loan in enumerate(loans, start=2):
                values = [
                    loan.loan_reference,
                    _borrower_name(loan),
                    float(loan.principal_amount or 0),
                    float(loan.balance or 0),
                    float(loan.installment_amount or 0),
                    float(loan.amount_paid or 0),
                    _date_value(loan.first_payment_due),
                    _date_value(loan.maturity_date),
                    _enum_value(loan.status),
                    _enum_value(loan.risk_level),
                    bool(loan.is_overdue),
                    loan.branch.name if loan.branch else "",
                ]
                for column, value in enumerate(values, start=1):
                    style = (
                        {"number_format": "M #,##0.00"}
                        if column in {3, 4, 5, 6}
                        else {}
                    )
                    cells[f"{get_column_letter(column)}{row}"] = {
                        "value": value,
                        **({"style": style} if style else {}),
                    }
            sheet.update(
                {
                    "cells": cells,
                    "row_count": max(40, len(loans) + 5),
                    "column_count": len(headers),
                    "freeze_panes": "A2",
                    "column_widths": {
                        "A": 21,
                        "B": 25,
                        "C": 16,
                        "D": 16,
                        "E": 16,
                        "F": 16,
                        "G": 15,
                        "H": 15,
                        "I": 14,
                        "J": 12,
                        "K": 12,
                        "L": 18,
                    },
                }
            )

    elif template == "cashbook":
        payment_query = (
            db.query(PaymentTransaction)
            .options(
                joinedload(PaymentTransaction.loan),
                joinedload(PaymentTransaction.borrower)
                .joinedload(Borrower.user)
                .joinedload(User.person),
            )
            .filter(PaymentTransaction.company_id == context.company_id)
        )
        if branch_id is not None:
            payment_query = payment_query.filter(
                PaymentTransaction.loan.has(ClientCompanyLoan.branch_id == branch_id)
            )
        payments = (
            payment_query.order_by(PaymentTransaction.created_at.desc())
            .limit(3000)
            .all()
        )
        headers = [
            "Date",
            "Reference",
            "Borrower",
            "Loan",
            "Purpose",
            "Method",
            "Direction",
            "Amount",
            "Currency",
            "Status",
        ]
        cells = _header_cells(headers, "#ECFDF3")
        for row, payment in enumerate(reversed(payments), start=2):
            amount = float(payment.amount or 0)
            direction = _enum_value(payment.direction)
            borrower = payment.borrower
            person = borrower.user.person if borrower and borrower.user else None
            values = [
                _date_value(payment.completed_at or payment.created_at),
                payment.provider_reference or payment.proof_reference or str(payment.id),
                person.full_name if person else "",
                payment.loan.loan_reference if payment.loan else "",
                _enum_value(payment.purpose),
                _enum_value(payment.payment_method),
                direction,
                amount,
                payment.currency,
                _enum_value(payment.status),
            ]
            for column, value in enumerate(values, start=1):
                style = {"number_format": "M #,##0.00"} if column == 8 else {}
                cells[f"{get_column_letter(column)}{row}"] = {
                    "value": value,
                    **({"style": style} if style else {}),
                }
        sheet.update(
            {
                "cells": cells,
                "row_count": max(40, len(payments) + 5),
                "column_count": len(headers),
                "freeze_panes": "A2",
                "column_widths": {
                    "A": 20,
                    "B": 24,
                    "C": 24,
                    "D": 20,
                    "E": 18,
                    "F": 15,
                    "G": 14,
                    "H": 16,
                    "I": 10,
                    "J": 14,
                },
            }
        )
    return workbook


def _is_spreadsheet(document: WorkspaceDocument) -> bool:
    return str(document.template_key or "").startswith("spreadsheet_")


def _spreadsheet_or_404(
    db: Session,
    context: TenantContext,
    document_id: UUID,
) -> WorkspaceDocument:
    document = _governed_document_or_404(db, context, document_id)
    if not _is_spreadsheet(document):
        raise HTTPException(status_code=404, detail="Spreadsheet not found")
    return document


def _safe_title(title: str | None, fallback: str = "Untitled workbook") -> str:
    return (title or fallback).strip()[:255] or fallback


def _create_document(
    db: Session,
    context: TenantContext,
    *,
    title: str,
    workbook: dict[str, Any],
    visibility: str,
    is_confidential: bool,
    template_key: str,
) -> WorkspaceDocument:
    _validate_visibility(context, visibility)
    normalized = normalize_workbook(workbook)
    document = WorkspaceDocument(
        owner_user_id=context.user.id,
        company_id=context.company_id,
        branch_id=context.branch_id,
        last_edited_by_user_id=context.user.id,
        reference=document_reference(),
        title=_safe_title(title),
        template_key=template_key,
        content_json=normalized,
        content_html="",
        plain_text=workbook_plain_text(normalized),
        visibility=visibility,
        status="draft",
        version=1,
        style_key="minimal_clean",
        default_font_family="Arial",
        default_font_size_pt=10,
        default_line_height_percent=100,
        include_brand_header=False,
        include_footer=False,
        is_confidential=is_confidential,
        cover_page_enabled=False,
        cover_page={},
        address_blocks=[],
    )
    _commit_created_document(db, document, user_id=context.user.id)
    return _governed_document_or_404(db, context, document.id)


@router.post("", response_model=WorkspaceDocumentRead, status_code=status.HTTP_201_CREATED)
def create_workspace_spreadsheet(
    payload: WorkspaceSpreadsheetCreate,
    db: Session = Depends(get_db),
    context: TenantContext = Depends(get_user_context),
):
    workbook = _company_data_workbook(db, context, payload.template)
    document = _create_document(
        db,
        context,
        title=payload.title,
        workbook=workbook,
        visibility=payload.visibility,
        is_confidential=payload.is_confidential,
        template_key=f"spreadsheet_{payload.template}",
    )
    return _read(db, context, document)


@router.post("/import", response_model=WorkspaceDocumentRead, status_code=status.HTTP_201_CREATED)
async def import_workspace_spreadsheet(
    file: UploadFile = File(...),
    title: str | None = Form(default=None),
    visibility: SpreadsheetVisibility = Form(default="private"),
    is_confidential: bool = Form(default=False),
    db: Session = Depends(get_db),
    context: TenantContext = Depends(get_user_context),
):
    filename = file.filename or "workbook.xlsx"
    try:
        workbook = import_workbook(await file.read(), filename)
    except SpreadsheetValidationError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    document = _create_document(
        db,
        context,
        title=_safe_title(title, Path(filename).stem or "Imported workbook"),
        workbook=workbook,
        visibility=visibility,
        is_confidential=is_confidential,
        template_key="spreadsheet_imported",
    )
    return _read(db, context, document)


@router.patch("/{document_id}", response_model=WorkspaceDocumentRead)
def update_workspace_spreadsheet(
    document_id: UUID,
    payload: WorkspaceSpreadsheetUpdate,
    db: Session = Depends(get_db),
    context: TenantContext = Depends(get_user_context),
):
    document = _spreadsheet_or_404(db, context, document_id)
    if not _can_edit(db, context, document):
        raise HTTPException(status_code=403, detail="You have view-only access to this spreadsheet")
    if document.status == "archived":
        raise HTTPException(status_code=409, detail="Archived spreadsheets cannot be edited")
    if payload.expected_version is not None and payload.expected_version != document.version:
        raise HTTPException(
            status_code=409,
            detail={
                "message": "This spreadsheet was changed by another user. Reload before saving.",
                "current_version": document.version,
            },
        )
    try:
        normalized = normalize_workbook(payload.workbook)
    except SpreadsheetValidationError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    if payload.title is not None:
        document.title = _safe_title(payload.title)
    document.content_json = normalized
    document.plain_text = workbook_plain_text(normalized)
    document.content_html = ""
    document.version += 1
    document.last_edited_by_user_id = context.user.id
    if payload.create_revision:
        _revision(db, document, context.user.id)
    db.commit()
    db.refresh(document)
    return _read(db, context, _spreadsheet_or_404(db, context, document.id))


@router.post("/{document_id}/refresh-loanhub-data", response_model=WorkspaceDocumentRead)
def refresh_workspace_spreadsheet_data(
    document_id: UUID,
    db: Session = Depends(get_db),
    context: TenantContext = Depends(get_user_context),
):
    document = _spreadsheet_or_404(db, context, document_id)
    if not _can_edit(db, context, document):
        raise HTTPException(status_code=403, detail="You have view-only access to this spreadsheet")
    if document.status == "archived":
        raise HTTPException(status_code=409, detail="Archived spreadsheets cannot be refreshed")
    template = str(document.template_key or "").removeprefix("spreadsheet_")
    if template not in DATABASE_TEMPLATES:
        raise HTTPException(status_code=400, detail="This workbook is not linked to a LoanHub database template")
    workbook = normalize_workbook(_company_data_workbook(db, context, template))
    document.content_json = workbook
    document.plain_text = workbook_plain_text(workbook)
    document.version += 1
    document.last_edited_by_user_id = context.user.id
    _revision(db, document, context.user.id)
    db.commit()
    db.refresh(document)
    return _read(db, context, _spreadsheet_or_404(db, context, document.id))


def _export_payload(
    document: WorkspaceDocument,
    format_name: SpreadsheetExportFormat,
) -> tuple[bytes, str, str]:
    try:
        if format_name == "xlsx":
            return (
                export_xlsx(document.content_json or {}),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "xlsx",
            )
        return export_csv(document.content_json or {}), "text/csv; charset=utf-8", "csv"
    except SpreadsheetValidationError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@router.get("/{document_id}/export/{format_name}")
def download_workspace_spreadsheet(
    document_id: UUID,
    format_name: SpreadsheetExportFormat,
    db: Session = Depends(get_db),
    context: TenantContext = Depends(get_user_context),
):
    document = _spreadsheet_or_404(db, context, document_id)
    content, mime_type, extension = _export_payload(document, format_name)
    clean_name = "".join(
        character if character.isalnum() or character in " -_" else "_"
        for character in document.title
    ).strip() or "workbook"
    filename = f"{clean_name}.{extension}"
    encoded = quote(filename, safe="")
    return StreamingResponse(
        BytesIO(content),
        media_type=mime_type,
        headers={
            "Content-Disposition": f"attachment; filename=\"workbook.{extension}\"; filename*=UTF-8''{encoded}",
            "Content-Length": str(len(content)),
            "Cache-Control": "private, no-store",
            "X-Content-Type-Options": "nosniff",
        },
    )


@router.post(
    "/{document_id}/publish",
    response_model=ManagedFileRead,
    status_code=status.HTTP_201_CREATED,
)
def publish_workspace_spreadsheet(
    document_id: UUID,
    payload: WorkspaceSpreadsheetPublish,
    db: Session = Depends(get_db),
    context: TenantContext = Depends(get_user_context),
):
    document = _spreadsheet_or_404(db, context, document_id)
    if not _can_edit(db, context, document):
        raise HTTPException(status_code=403, detail="You have view-only access to this spreadsheet")
    content, mime_type, extension = _export_payload(document, payload.format)
    visibility = payload.visibility or ("company" if document.company_id else "private")
    _validate_visibility(context, visibility)
    clean_title = "".join(
        character if character.isalnum() or character in " -_" else "_"
        for character in document.title
    ).strip() or "workbook"
    name = payload.file_name or f"{clean_title}.{extension}"
    if not name.lower().endswith(f".{extension}"):
        name = f"{name}.{extension}"
    record = store_bytes(
        db,
        content=content,
        original_name=name,
        mime_type=mime_type,
        owner_user_id=context.user.id,
        company_id=document.company_id,
        branch_id=document.branch_id,
        category="spreadsheet_document",
        visibility=visibility,
        description=f"Published from {document.reference}, version {document.version}",
        linked_entity_type="workspace_spreadsheet",
        linked_entity_id=str(document.id),
        is_confidential=(
            document.is_confidential
            if payload.is_confidential is None
            else payload.is_confidential
        ),
    )
    db.commit()
    db.refresh(record)
    return record
