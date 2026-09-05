from __future__ import annotations

import csv
from datetime import date, datetime
from io import BytesIO, StringIO
from pathlib import Path
import re
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, range_boundaries


MAX_UPLOAD_BYTES = 12 * 1024 * 1024
MAX_SHEETS = 30
MAX_ROWS = 5000
MAX_COLUMNS = 200
MAX_CELLS = 100_000
MAX_TEXT_LENGTH = 50_000
CELL_REFERENCE = re.compile(r"^[A-Z]{1,3}[1-9][0-9]*$")


class SpreadsheetValidationError(ValueError):
    pass


def _safe_scalar(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def _color(value: Any) -> str | None:
    rgb = getattr(value, "rgb", None)
    if isinstance(rgb, str) and len(rgb) >= 6:
        return f"#{rgb[-6:]}"
    return None


def _side_to_json(side: Side) -> dict[str, Any] | None:
    if not side or not side.style:
        return None
    return {"style": side.style, "color": _color(side.color)}


def _cell_style(cell: Any) -> dict[str, Any]:
    style: dict[str, Any] = {}
    font = cell.font
    if font:
        if font.bold:
            style["bold"] = True
        if font.italic:
            style["italic"] = True
        if font.underline:
            style["underline"] = True
        if font.size:
            style["font_size"] = float(font.size)
        color = _color(font.color)
        if color:
            style["font_color"] = color
    fill = cell.fill
    fill_color = _color(getattr(fill, "fgColor", None))
    if fill_color and fill_color.upper() not in {"#000000", "#FFFFFF"}:
        style["fill_color"] = fill_color
    alignment = cell.alignment
    if alignment:
        if alignment.horizontal:
            style["horizontal"] = alignment.horizontal
        if alignment.vertical:
            style["vertical"] = alignment.vertical
        if alignment.wrap_text:
            style["wrap"] = True
    if cell.number_format and cell.number_format != "General":
        style["number_format"] = cell.number_format
    border = cell.border
    if border:
        border_json = {
            name: item
            for name in ("left", "right", "top", "bottom")
            if (item := _side_to_json(getattr(border, name))) is not None
        }
        if border_json:
            style["border"] = border_json
    return style


def blank_workbook(*, template: str = "blank") -> dict[str, Any]:
    cells: dict[str, dict[str, Any]] = {}
    row_count = 40
    column_count = 12
    if template == "loan_portfolio":
        headers = ["Loan reference", "Borrower", "Principal", "Balance", "Installment", "Next due", "Status"]
        for index, header in enumerate(headers, start=1):
            cells[f"{get_column_letter(index)}1"] = {
                "value": header,
                "style": {"bold": True, "fill_color": "#EAF2FF", "horizontal": "center"},
            }
        column_count = len(headers)
    elif template == "cashbook":
        headers = ["Date", "Reference", "Description", "Money in", "Money out", "Balance", "Notes"]
        for index, header in enumerate(headers, start=1):
            cells[f"{get_column_letter(index)}1"] = {
                "value": header,
                "style": {"bold": True, "fill_color": "#ECFDF3", "horizontal": "center"},
            }
        column_count = len(headers)
    elif template == "collections":
        headers = ["Borrower", "Loan", "Amount due", "Days overdue", "Last contact", "Promise date", "Outcome", "Owner"]
        for index, header in enumerate(headers, start=1):
            cells[f"{get_column_letter(index)}1"] = {
                "value": header,
                "style": {"bold": True, "fill_color": "#FFF4E5", "horizontal": "center"},
            }
        column_count = len(headers)
    elif template == "budget":
        headers = ["Category", "Budget", "Actual", "Variance", "Variance %", "Notes"]
        for index, header in enumerate(headers, start=1):
            cells[f"{get_column_letter(index)}1"] = {
                "value": header,
                "style": {"bold": True, "fill_color": "#F3E8FF", "horizontal": "center"},
            }
        cells["D2"] = {"formula": "=B2-C2", "value": "=B2-C2"}
        cells["E2"] = {"formula": "=IF(B2=0,0,D2/B2)", "value": "=IF(B2=0,0,D2/B2)", "style": {"number_format": "0.00%"}}
        column_count = len(headers)
    return {
        "format_version": 1,
        "active_sheet": "Sheet1",
        "sheets": [
            {
                "id": "sheet-1",
                "name": "Sheet1",
                "row_count": row_count,
                "column_count": column_count,
                "cells": cells,
                "column_widths": {},
                "row_heights": {},
                "merges": [],
                "freeze_panes": "A2" if cells else None,
            }
        ],
    }


def normalize_workbook(payload: Any) -> dict[str, Any]:
    if not isinstance(payload, dict):
        raise SpreadsheetValidationError("Workbook data must be an object")
    source_sheets = payload.get("sheets")
    if not isinstance(source_sheets, list) or not source_sheets:
        raise SpreadsheetValidationError("A workbook must contain at least one sheet")
    if len(source_sheets) > MAX_SHEETS:
        raise SpreadsheetValidationError(f"A workbook can contain at most {MAX_SHEETS} sheets")

    sheets: list[dict[str, Any]] = []
    seen_names: set[str] = set()
    total_cells = 0
    for sheet_index, raw_sheet in enumerate(source_sheets, start=1):
        if not isinstance(raw_sheet, dict):
            raise SpreadsheetValidationError("Every sheet must be an object")
        name = str(raw_sheet.get("name") or f"Sheet{sheet_index}").strip()[:31] or f"Sheet{sheet_index}"
        base_name = name
        suffix = 2
        while name.lower() in seen_names:
            name = f"{base_name[:27]} {suffix}"[:31]
            suffix += 1
        seen_names.add(name.lower())

        try:
            row_count = int(raw_sheet.get("row_count") or 40)
            column_count = int(raw_sheet.get("column_count") or 12)
        except (TypeError, ValueError) as exc:
            raise SpreadsheetValidationError("Sheet dimensions must be numbers") from exc
        row_count = max(1, min(row_count, MAX_ROWS))
        column_count = max(1, min(column_count, MAX_COLUMNS))

        raw_cells = raw_sheet.get("cells") or {}
        if not isinstance(raw_cells, dict):
            raise SpreadsheetValidationError("Sheet cells must be an object")
        cells: dict[str, dict[str, Any]] = {}
        for reference, raw_cell in raw_cells.items():
            ref = str(reference).upper().strip()
            if not CELL_REFERENCE.fullmatch(ref):
                continue
            if not isinstance(raw_cell, dict):
                raw_cell = {"value": raw_cell}
            min_col, min_row, max_col, max_row = range_boundaries(ref)
            if max_row > MAX_ROWS or max_col > MAX_COLUMNS:
                raise SpreadsheetValidationError(
                    f"Cell {ref} exceeds the supported workbook size ({MAX_ROWS} rows x {MAX_COLUMNS} columns)"
                )
            row_count = max(row_count, max_row)
            column_count = max(column_count, max_col)
            value = _safe_scalar(raw_cell.get("value"))
            formula = raw_cell.get("formula")
            if formula is not None:
                formula = str(formula).strip()
                if formula and not formula.startswith("="):
                    formula = f"={formula}"
                if len(formula) > 8192:
                    raise SpreadsheetValidationError(f"Formula in {ref} is too long")
            elif isinstance(value, str) and value.startswith("="):
                formula = value
            style = raw_cell.get("style") if isinstance(raw_cell.get("style"), dict) else {}
            cells[ref] = {
                "value": value,
                **({"formula": formula} if formula else {}),
                **({"style": style} if style else {}),
                **({"value_type": str(raw_cell.get("value_type"))} if raw_cell.get("value_type") else {}),
            }
            total_cells += 1
            if total_cells > MAX_CELLS:
                raise SpreadsheetValidationError(f"A workbook can contain at most {MAX_CELLS:,} populated cells")

        merges = []
        for raw_range in raw_sheet.get("merges") or []:
            token = str(raw_range).upper().strip()
            try:
                min_col, min_row, max_col, max_row = range_boundaries(token)
            except ValueError:
                continue
            if max_row <= MAX_ROWS and max_col <= MAX_COLUMNS and (min_row != max_row or min_col != max_col):
                merges.append(token)

        column_widths: dict[str, float] = {}
        for key, value in (raw_sheet.get("column_widths") or {}).items():
            try:
                width = float(value)
            except (TypeError, ValueError):
                continue
            column_widths[str(key).upper()] = max(2.0, min(width, 100.0))
        row_heights: dict[str, float] = {}
        for key, value in (raw_sheet.get("row_heights") or {}).items():
            try:
                height = float(value)
            except (TypeError, ValueError):
                continue
            row_heights[str(key)] = max(8.0, min(height, 200.0))

        freeze_panes = raw_sheet.get("freeze_panes")
        if freeze_panes is not None:
            freeze_panes = str(freeze_panes).upper().strip() or None
            if freeze_panes and not CELL_REFERENCE.fullmatch(freeze_panes):
                freeze_panes = None

        sheets.append(
            {
                "id": str(raw_sheet.get("id") or f"sheet-{sheet_index}")[:80],
                "name": name,
                "row_count": row_count,
                "column_count": column_count,
                "cells": cells,
                "column_widths": column_widths,
                "row_heights": row_heights,
                "merges": merges,
                "freeze_panes": freeze_panes,
            }
        )

    active_sheet = str(payload.get("active_sheet") or sheets[0]["name"])
    if active_sheet not in {sheet["name"] for sheet in sheets}:
        active_sheet = sheets[0]["name"]
    return {
        "format_version": 1,
        "active_sheet": active_sheet,
        "sheets": sheets,
        **({"source_filename": str(payload.get("source_filename"))[:255]} if payload.get("source_filename") else {}),
    }


def _worksheet_to_json(worksheet: Any, sheet_index: int) -> dict[str, Any]:
    if worksheet.max_row > MAX_ROWS or worksheet.max_column > MAX_COLUMNS:
        raise SpreadsheetValidationError(
            f"Sheet '{worksheet.title}' is larger than {MAX_ROWS} rows x {MAX_COLUMNS} columns"
        )
    cells: dict[str, dict[str, Any]] = {}
    populated = 0
    for row in worksheet.iter_rows():
        for cell in row:
            if isinstance(cell, MergedCell) or cell.value is None:
                continue
            populated += 1
            if populated > MAX_CELLS:
                raise SpreadsheetValidationError(f"Sheet '{worksheet.title}' contains too many populated cells")
            raw_value = cell.value
            value_type = None
            formula = None
            if isinstance(raw_value, str) and raw_value.startswith("="):
                formula = raw_value
            if isinstance(raw_value, (datetime, date)):
                value_type = "date"
            entry: dict[str, Any] = {"value": _safe_scalar(raw_value)}
            if formula:
                entry["formula"] = formula
            style = _cell_style(cell)
            if style:
                entry["style"] = style
            if value_type:
                entry["value_type"] = value_type
            cells[cell.coordinate] = entry

    column_widths = {
        key: float(dimension.width)
        for key, dimension in worksheet.column_dimensions.items()
        if dimension.width is not None
    }
    row_heights = {
        str(key): float(dimension.height)
        for key, dimension in worksheet.row_dimensions.items()
        if dimension.height is not None
    }
    return {
        "id": f"sheet-{sheet_index}",
        "name": worksheet.title,
        "row_count": max(40, min(max(worksheet.max_row, 1), MAX_ROWS)),
        "column_count": max(12, min(max(worksheet.max_column, 1), MAX_COLUMNS)),
        "cells": cells,
        "column_widths": column_widths,
        "row_heights": row_heights,
        "merges": [str(item) for item in worksheet.merged_cells.ranges],
        "freeze_panes": str(worksheet.freeze_panes) if worksheet.freeze_panes else None,
    }


def import_workbook(content: bytes, filename: str) -> dict[str, Any]:
    if not content:
        raise SpreadsheetValidationError("The uploaded spreadsheet is empty")
    if len(content) > MAX_UPLOAD_BYTES:
        raise SpreadsheetValidationError(f"Spreadsheet uploads are limited to {MAX_UPLOAD_BYTES // 1024 // 1024} MB")
    suffix = Path(filename or "").suffix.lower()
    if suffix in {".csv", ".tsv"}:
        encoding = "utf-8-sig"
        try:
            text = content.decode(encoding)
        except UnicodeDecodeError:
            text = content.decode("latin-1")
        delimiter = "\t" if suffix == ".tsv" else ","
        rows = list(csv.reader(StringIO(text), delimiter=delimiter))
        if len(rows) > MAX_ROWS:
            raise SpreadsheetValidationError(f"Delimited files are limited to {MAX_ROWS} rows")
        cells: dict[str, dict[str, Any]] = {}
        max_columns = 1
        for row_index, row in enumerate(rows, start=1):
            if len(row) > MAX_COLUMNS:
                raise SpreadsheetValidationError(f"Delimited files are limited to {MAX_COLUMNS} columns")
            max_columns = max(max_columns, len(row))
            for column_index, value in enumerate(row, start=1):
                if value != "":
                    cells[f"{get_column_letter(column_index)}{row_index}"] = {"value": value}
        workbook = blank_workbook()
        sheet = workbook["sheets"][0]
        sheet["name"] = Path(filename).stem[:31] or "Sheet1"
        sheet["row_count"] = max(40, len(rows))
        sheet["column_count"] = max(12, max_columns)
        sheet["cells"] = cells
        workbook["active_sheet"] = sheet["name"]
        workbook["source_filename"] = filename
        return normalize_workbook(workbook)

    if suffix not in {".xlsx", ".xlsm"}:
        raise SpreadsheetValidationError("Upload an .xlsx, .xlsm, .csv or .tsv spreadsheet")
    try:
        source = load_workbook(BytesIO(content), data_only=False, read_only=False, keep_vba=False)
    except Exception as exc:
        raise SpreadsheetValidationError("The Excel workbook could not be opened") from exc
    if len(source.worksheets) > MAX_SHEETS:
        raise SpreadsheetValidationError(f"A workbook can contain at most {MAX_SHEETS} sheets")
    sheets = [_worksheet_to_json(sheet, index) for index, sheet in enumerate(source.worksheets, start=1)]
    return normalize_workbook(
        {
            "format_version": 1,
            "active_sheet": source.active.title if source.worksheets else sheets[0]["name"],
            "sheets": sheets,
            "source_filename": filename,
        }
    )


def _parse_excel_value(entry: dict[str, Any]) -> Any:
    formula = entry.get("formula")
    if formula:
        return str(formula)
    value = entry.get("value")
    if entry.get("value_type") == "date" and isinstance(value, str):
        try:
            return datetime.fromisoformat(value)
        except ValueError:
            return value
    return value


def _apply_style(cell: Any, style: dict[str, Any]) -> None:
    if not style:
        return
    font_color = str(style.get("font_color") or "").replace("#", "") or None
    cell.font = Font(
        name=cell.font.name,
        size=float(style.get("font_size") or cell.font.sz or 11),
        bold=bool(style.get("bold")),
        italic=bool(style.get("italic")),
        underline="single" if style.get("underline") else None,
        color=font_color,
    )
    fill_color = str(style.get("fill_color") or "").replace("#", "")
    if fill_color:
        cell.fill = PatternFill(fill_type="solid", fgColor=fill_color)
    cell.alignment = Alignment(
        horizontal=style.get("horizontal") or cell.alignment.horizontal,
        vertical=style.get("vertical") or cell.alignment.vertical,
        wrap_text=bool(style.get("wrap")),
    )
    if style.get("number_format"):
        cell.number_format = str(style["number_format"])
    border_payload = style.get("border") if isinstance(style.get("border"), dict) else {}
    if border_payload:
        def make_side(name: str) -> Side:
            item = border_payload.get(name) if isinstance(border_payload.get(name), dict) else {}
            color = str(item.get("color") or "").replace("#", "") or None
            return Side(style=item.get("style"), color=color)
        cell.border = Border(
            left=make_side("left"), right=make_side("right"),
            top=make_side("top"), bottom=make_side("bottom"),
        )


def export_xlsx(payload: Any) -> bytes:
    workbook_payload = normalize_workbook(payload)
    target = Workbook()
    target.remove(target.active)
    for sheet_payload in workbook_payload["sheets"]:
        worksheet = target.create_sheet(title=sheet_payload["name"])
        for reference, entry in sheet_payload["cells"].items():
            cell = worksheet[reference]
            cell.value = _parse_excel_value(entry)
            _apply_style(cell, entry.get("style") or {})
        for column, width in sheet_payload.get("column_widths", {}).items():
            worksheet.column_dimensions[column].width = width
        for row, height in sheet_payload.get("row_heights", {}).items():
            worksheet.row_dimensions[int(row)].height = height
        for merged_range in sheet_payload.get("merges", []):
            try:
                worksheet.merge_cells(merged_range)
            except ValueError:
                continue
        if sheet_payload.get("freeze_panes"):
            worksheet.freeze_panes = sheet_payload["freeze_panes"]
    for index, worksheet in enumerate(target.worksheets):
        if worksheet.title == workbook_payload["active_sheet"]:
            target.active = index
            break
    output = BytesIO()
    target.save(output)
    return output.getvalue()


def export_csv(payload: Any) -> bytes:
    workbook_payload = normalize_workbook(payload)
    active_name = workbook_payload["active_sheet"]
    sheet = next((item for item in workbook_payload["sheets"] if item["name"] == active_name), workbook_payload["sheets"][0])
    output = StringIO(newline="")
    writer = csv.writer(output)
    row_limit = max(1, int(sheet["row_count"]))
    column_limit = max(1, int(sheet["column_count"]))
    for row_index in range(1, row_limit + 1):
        row: list[Any] = []
        for column_index in range(1, column_limit + 1):
            entry = sheet["cells"].get(f"{get_column_letter(column_index)}{row_index}", {})
            row.append(entry.get("formula") or entry.get("value") or "")
        while row and row[-1] == "":
            row.pop()
        writer.writerow(row)
    return output.getvalue().encode("utf-8-sig")


def workbook_plain_text(payload: Any) -> str:
    workbook_payload = normalize_workbook(payload)
    chunks: list[str] = []
    remaining = MAX_TEXT_LENGTH
    for sheet in workbook_payload["sheets"]:
        chunks.append(f"[{sheet['name']}]")
        remaining -= len(chunks[-1])
        for reference, entry in sheet["cells"].items():
            value = entry.get("formula") or entry.get("value")
            if value in (None, ""):
                continue
            line = f"{reference}: {value}"
            if len(line) > remaining:
                return "\n".join(chunks)[:MAX_TEXT_LENGTH]
            chunks.append(line)
            remaining -= len(line) + 1
    return "\n".join(chunks)[:MAX_TEXT_LENGTH]
