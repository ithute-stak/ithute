from inspect import signature
from io import BytesIO

import pytest
from docx import Document as WordDocument
from pydantic import ValidationError

from database.schemas.company_website import CompanyWebsiteUpdate
from routers.notifications import list_notifications, unread_count
from routers.performance_system_ratings import _evidence_level, _rating
from routers.workspace_document_organization import _docx_html, _normalise_folder_name


def test_notification_surfaces_default_to_attention_only():
    list_params = signature(list_notifications).parameters
    count_params = signature(unread_count).parameters

    assert list_params["attention_only"].default is True
    assert list_params["include_routine"].default is False
    assert count_params["include_routine"].default is False


def test_company_custom_sections_reject_script_style_content():
    with pytest.raises(ValidationError):
        CompanyWebsiteUpdate(custom_sections=[{"script": "alert('x')"}])

    with pytest.raises(ValidationError):
        CompanyWebsiteUpdate(custom_sections=[{"button_url": "javascript:alert(1)"}])


def test_company_custom_sections_are_bounded():
    with pytest.raises(ValidationError):
        CompanyWebsiteUpdate(custom_sections=[{"title": f"Section {index}"} for index in range(13)])


def test_document_folder_names_are_cleaned():
    assert _normalise_folder_name("  Credit   committee  ") == ("Credit committee", "credit committee")


def test_word_import_extracts_editable_text_headings_and_tables():
    source = WordDocument()
    source.add_heading("Collections report", level=1)
    source.add_paragraph("This paragraph must remain editable.")
    table = source.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Borrower"
    table.cell(0, 1).text = "Balance"
    stream = BytesIO()
    source.save(stream)

    imported = _docx_html(stream.getvalue())

    assert "<h1>Collections report</h1>" in imported
    assert "This paragraph must remain editable." in imported
    assert "<table>" in imported
    assert "Borrower" in imported
    assert "Balance" in imported


def test_system_rating_bands_and_confidence_are_deterministic():
    assert _rating(91) == "Exceptional"
    assert _rating(82) == "Exceeds expectations"
    assert _rating(61) == "Meets expectations"
    assert _rating(49) == "Needs attention"

    assert _evidence_level(80) == "high"
    assert _evidence_level(60) == "moderate"
    assert _evidence_level(30) == "limited"
    assert _evidence_level(0) == "insufficient"
