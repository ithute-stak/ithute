from __future__ import annotations

import hashlib
import re
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any, Literal

from fastapi import APIRouter, Depends, HTTPException, Query
from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.api.v1.procurement import commit, policy as procurement_policy, require_anywhere, require_scope, row_dict
from app.core.config import get_settings
from app.db.session import get_db
from app.models import (
    AlgorithmicAssistantAnalysis,
    Document,
    DocumentVersion,
    ProcurementRequisition,
    ProcurementRequisitionLine,
    PurchaseOrder,
    PurchaseOrderLine,
    Supplier,
    SupplierQuotation,
    SupplierQuotationLine,
    Tender,
    TenderChecklistItem,
)
from app.security.access import Principal, current_principal


router = APIRouter(prefix="/assistants", tags=["Algorithmic Procurement and Tender Assistants"])
RULE_VERSION = "algorithmic-v1"

CATEGORY_RULES: dict[str, tuple[str, ...]] = {
    "materials": ("cement", "sand", "stone", "aggregate", "brick", "steel", "timber", "concrete", "pipe", "paint"),
    "fuel": ("diesel", "petrol", "fuel", "lubricant", "oil"),
    "plant_and_tools": ("excavator", "grader", "generator", "plant", "tool", "drill", "compactor"),
    "ppe_and_safety": ("ppe", "helmet", "boot", "glove", "safety", "vest", "harness"),
    "professional_services": ("consulting", "survey", "testing", "legal", "audit", "design"),
    "office_and_it": ("laptop", "printer", "stationery", "software", "internet", "office"),
    "subcontract_services": ("subcontract", "labour-only", "labour only", "specialist works"),
}

TENDER_REQUIREMENT_RULES: tuple[tuple[str, str, tuple[str, ...]], ...] = (
    ("company", "Tax clearance", ("tax clearance", "tax compliance", "tax certificate")),
    ("company", "Company registration", ("company registration", "certificate of incorporation", "registration certificate")),
    ("company", "Financial statements", ("financial statement", "audited account", "bank statement")),
    ("technical", "Previous experience", ("previous experience", "similar project", "track record", "past experience")),
    ("technical", "Method statement", ("method statement", "methodology", "technical proposal")),
    ("commercial", "Bill of quantities (BOQ)", ("bill of quantities", "boq", "pricing schedule")),
    ("commercial", "Bid security", ("bid security", "bid bond", "tender security", "bank guarantee")),
    ("company", "Insurance", ("insurance", "insurance certificate", "public liability")),
    ("technical", "Health and safety documents", ("health and safety", "hse plan", "safety plan")),
)


class SourceInput(BaseModel):
    model_config = ConfigDict(extra="forbid")
    document_id: int | None = None
    source_text: str | None = Field(default=None, max_length=800_000)


class DraftInput(BaseModel):
    model_config = ConfigDict(extra="forbid")
    kind: Literal["company_profile", "technical_response", "method_statement", "cover_letter"]
    document_id: int | None = None


def serialise(value: Any) -> Any:
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, dict):
        return {str(key): serialise(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [serialise(item) for item in value]
    return value


def as_utc(value: datetime) -> datetime:
    return value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value.astimezone(timezone.utc)


def sentences(text: str) -> list[str]:
    return [part.strip() for part in re.split(r"(?<=[.!?;])\s+|\n+", re.sub(r"\s+", " ", text)) if part.strip()]


def matching_snippets(text: str, terms: tuple[str, ...] | list[str], limit: int = 4) -> list[str]:
    found: list[str] = []
    for sentence in sentences(text):
        if any(term.casefold() in sentence.casefold() for term in terms):
            found.append(sentence[:500])
        if len(found) >= limit:
            break
    return found


def category_suggestion(*parts: str | None) -> dict[str, Any]:
    text = " ".join(part or "" for part in parts).casefold()
    scores = {category: sum(text.count(term) for term in terms) for category, terms in CATEGORY_RULES.items()}
    ranked = [{"category": category, "matched_terms": [term for term in terms if term in text], "score": score} for category, terms in CATEGORY_RULES.items() if (score := scores[category])]
    ranked.sort(key=lambda row: (-int(row["score"]), str(row["category"])))
    return {"suggested_category": ranked[0]["category"] if ranked else "unclassified", "candidates": ranked, "rule": "keyword frequency; staff must confirm the category"}


def latest_document_text(db: Session, company_id: int, document_id: int) -> tuple[str, str | None, str | None]:
    document = db.get(Document, document_id)
    if not document or document.company_id != company_id or document.status != "active":
        raise HTTPException(status_code=422, detail="The selected source document is not an active company document")
    version = db.scalar(select(DocumentVersion).where(DocumentVersion.document_id == document.id).order_by(DocumentVersion.version_number.desc()).limit(1))
    if not version:
        raise HTTPException(status_code=422, detail="The selected source document has no uploaded version")
    media_root = Path(get_settings().media_root).resolve()
    path = (media_root / version.stored_path).resolve()
    try:
        path.relative_to(media_root)
    except ValueError as error:
        raise HTTPException(status_code=422, detail="The stored document path is invalid") from error
    if not path.is_file():
        raise HTTPException(status_code=422, detail="The selected document file is unavailable")
    suffix = path.suffix.casefold()
    try:
        if suffix == ".pdf":
            from pypdf import PdfReader
            text = "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
        elif suffix == ".docx":
            from docx import Document as WordDocument
            word = WordDocument(str(path))
            text = "\n".join(paragraph.text for paragraph in word.paragraphs)
            text += "\n" + "\n".join(" | ".join(cell.text for cell in row.cells) for table in word.tables for row in table.rows)
        elif suffix in {".txt", ".csv", ".md", ".json"}:
            text = path.read_text(encoding="utf-8", errors="replace")
        else:
            raise HTTPException(status_code=422, detail="Use PDF, DOCX, TXT, CSV, Markdown or JSON for deterministic document analysis")
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(status_code=422, detail="The document could not be read by the deterministic extractor") from error
    if not text.strip():
        raise HTTPException(status_code=422, detail="No readable text was found in the source document")
    return text, version.sha256, version.original_filename


def source_text(db: Session, principal: Principal, payload: SourceInput) -> tuple[str, int | None, str | None, str]:
    if payload.source_text and payload.source_text.strip():
        text = payload.source_text.strip()
        return text, payload.document_id, hashlib.sha256(text.encode("utf-8")).hexdigest(), "Pasted source text"
    if payload.document_id is None:
        raise HTTPException(status_code=422, detail="Provide pasted source text or select an uploaded document")
    text, digest, filename = latest_document_text(db, principal.user.company_id, payload.document_id)
    return text, payload.document_id, digest, filename or "Uploaded document"


def save_analysis(db: Session, principal: Principal, *, module: str, analysis_type: str, subject_type: str, subject_id: int | str, branch_id: int | None, site_id: int | None, document_id: int | None, source_sha256: str | None, result: dict[str, Any]) -> dict[str, Any]:
    row = AlgorithmicAssistantAnalysis(
        company_id=principal.user.company_id, branch_id=branch_id, site_id=site_id,
        module=module, analysis_type=analysis_type, subject_type=subject_type, subject_id=str(subject_id),
        document_id=document_id, source_sha256=source_sha256, rule_version=RULE_VERSION,
        result=serialise(result), created_by=principal.user.full_name,
    )
    db.add(row)
    db.flush()
    commit(db)
    return row_dict(row)


def requisition_or_404(db: Session, principal: Principal, requisition_id: int, permission: str = "procurement.view") -> ProcurementRequisition:
    row = db.get(ProcurementRequisition, requisition_id)
    if not row or row.company_id != principal.user.company_id:
        raise HTTPException(status_code=404, detail="Procurement requisition not found")
    require_scope(principal, permission, row.branch_id, row.site_id)
    return row


def tender_or_404(db: Session, principal: Principal, tender_id: int, permission: str = "tenders.view") -> Tender:
    row = db.get(Tender, tender_id)
    if not row or row.company_id != principal.user.company_id:
        raise HTTPException(status_code=404, detail="Tender not found")
    if not principal.can(permission, branch_id=row.branch_id, site_id=row.site_id):
        raise HTTPException(status_code=403, detail=f"Permission required in this branch/site: {permission}")
    return row


def review_requisition(row: ProcurementRequisition, lines: list[ProcurementRequisitionLine]) -> dict[str, Any]:
    findings: list[dict[str, str]] = []
    if len(row.title.strip()) < 6:
        findings.append({"severity": "blocking", "field": "title", "message": "Add a clear business purpose to the request title."})
    if not lines:
        findings.append({"severity": "blocking", "field": "lines", "message": "At least one requested item or service is required."})
    if row.required_by < date.today():
        findings.append({"severity": "blocking", "field": "required_by", "message": "The required-by date is in the past."})
    categories: list[dict[str, Any]] = []
    for line in lines:
        if not line.specification:
            findings.append({"severity": "warning", "field": f"line:{line.id}:specification", "message": f"Add a measurable specification for '{line.description}'."})
        if Decimal(line.estimated_unit_cost) <= 0:
            findings.append({"severity": "warning", "field": f"line:{line.id}:estimated_unit_cost", "message": f"Add an estimated unit cost for '{line.description}' to support the approval threshold."})
        categories.append({"line_id": line.id, "description": line.description, **category_suggestion(line.description, line.specification, line.notes)})
    overall = category_suggestion(row.title, row.notes, *(line.description for line in lines))
    return {
        "kind": "purchase_request_review", "recommendation_only": True,
        "requisition_number": row.requisition_number, "estimated_total": str(row.estimated_total),
        "findings": findings, "ready_for_submission": not any(item["severity"] == "blocking" for item in findings),
        "suggested_category": overall, "line_categories": categories,
        "explanation": "Rules check core request fields and score category keywords. They do not approve or submit the requisition.",
    }


def extract_terms(text: str) -> dict[str, Any]:
    clean = re.sub(r"\s+", " ", text)
    def period(label: str) -> int | None:
        match = re.search(label + r".{0,48}?(\d{1,4})\s*(day|days|month|months|year|years)", clean, flags=re.I)
        if not match:
            return None
        value = int(match.group(1)); unit = match.group(2).casefold()
        return value * (365 if unit.startswith("year") else 30 if unit.startswith("month") else 1)
    warranty = period(r"(?:warranty|guarantee)")
    payment = period(r"(?:payment\s+terms?|terms\s+of\s+payment)")
    delivery = period(r"(?:delivery|lead\s+time)")
    return {
        "warranty_days": warranty, "payment_terms_days_from_document": payment,
        "delivery_days_from_document": delivery,
        "evidence": matching_snippets(clean, ["warranty", "guarantee", "payment terms", "delivery", "lead time"], limit=8),
    }


def quote_comparison(db: Session, row: ProcurementRequisition) -> dict[str, Any]:
    quotes = db.scalars(select(SupplierQuotation).where(SupplierQuotation.requisition_id == row.id).order_by(SupplierQuotation.total_amount, SupplierQuotation.id)).all()
    if not quotes:
        raise HTTPException(status_code=409, detail="Record at least one supplier quotation before running a comparison")
    candidates: list[dict[str, Any]] = []
    for quote in quotes:
        supplier = db.get(Supplier, quote.supplier_id)
        terms: dict[str, Any] = {"warranty_days": None, "payment_terms_days_from_document": None, "delivery_days_from_document": None, "evidence": []}
        if quote.document_id:
            try:
                source, _, _ = latest_document_text(db, row.company_id, quote.document_id)
                terms = extract_terms(source)
            except HTTPException:
                terms["evidence"] = ["Attached document was not readable by the deterministic extractor."]
        candidates.append({
            "quotation_id": quote.id, "quote_reference": quote.quote_reference, "supplier_id": quote.supplier_id,
            "supplier": supplier.name if supplier else "Unknown supplier", "supplier_registration_number": supplier.registration_number if supplier else None,
            "supplier_contact": supplier.contact_name if supplier else None, "supplier_email": supplier.email if supplier else None,
            "total_amount": Decimal(quote.total_amount), "delivery_days": quote.delivery_days if quote.delivery_days is not None else terms["delivery_days_from_document"],
            "warranty_days": terms["warranty_days"], "payment_terms_days": terms["payment_terms_days_from_document"] if terms["payment_terms_days_from_document"] is not None else (supplier.payment_terms_days if supplier else None),
            "document_id": quote.document_id, "document_evidence": terms["evidence"], "status": quote.status,
        })
    price_floor = min(item["total_amount"] for item in candidates)
    known_delivery = [int(item["delivery_days"]) for item in candidates if item["delivery_days"] is not None]
    known_warranty = [int(item["warranty_days"]) for item in candidates if item["warranty_days"] is not None]
    known_payment = [int(item["payment_terms_days"]) for item in candidates if item["payment_terms_days"] is not None]
    for item in candidates:
        price_score = Decimal("50") if item["total_amount"] == 0 else (Decimal("50") * price_floor / item["total_amount"])
        delivery_score = Decimal("20") * min(known_delivery) / int(item["delivery_days"]) if item["delivery_days"] is not None and known_delivery else Decimal("0")
        warranty_score = Decimal("15") * int(item["warranty_days"]) / max(known_warranty) if item["warranty_days"] is not None and known_warranty else Decimal("0")
        payment_score = Decimal("15") * int(item["payment_terms_days"]) / max(known_payment) if item["payment_terms_days"] is not None and known_payment else Decimal("0")
        item["rule_scores"] = {"price": str(price_score.quantize(Decimal("0.01"))), "delivery": str(delivery_score.quantize(Decimal("0.01"))), "warranty": str(warranty_score.quantize(Decimal("0.01"))), "payment_terms": str(payment_score.quantize(Decimal("0.01")))}
        item["comparison_score"] = str((price_score + delivery_score + warranty_score + payment_score).quantize(Decimal("0.01")))
        item["total_amount"] = str(item["total_amount"])
    candidates.sort(key=lambda item: (-Decimal(item["comparison_score"]), Decimal(item["total_amount"]), str(item["supplier"])))
    return {
        "kind": "quotation_comparison", "recommendation_only": True,
        "requisition_number": row.requisition_number, "minimum_quotes_policy": procurement_policy(db, row.company_id)["minimum_quotes"],
        "weights": {"price": 50, "delivery": 20, "warranty": 15, "payment_terms": 15},
        "ranked_quotations": candidates,
        "recommendation": "The top-ranked quotation is a transparent commercial suggestion only. Staff must select the quotation and follow the maker/checker purchase-order workflow.",
    }


def extract_date_hits(text: str) -> list[dict[str, str]]:
    hits: list[dict[str, str]] = []
    for sentence in sentences(text):
        if not any(word in sentence.casefold() for word in ("deadline", "closing", "submission", "site visit", "clarification")):
            continue
        values = re.findall(r"\b(?:\d{4}-\d{2}-\d{2}|\d{1,2}[/-]\d{1,2}[/-]\d{2,4})\b", sentence)
        for value in values:
            hits.append({"value": value, "evidence": sentence[:500]})
    return hits[:12]


def tender_document_review(text: str, tender: Tender) -> dict[str, Any]:
    required = []
    for category, name, terms in TENDER_REQUIREMENT_RULES:
        evidence = matching_snippets(text, list(terms), limit=2)
        if evidence:
            required.append({"category": category, "name": name, "required": True, "evidence": evidence})
    technical = [sentence[:500] for sentence in sentences(text) if re.search(r"\b(must|shall|required|minimum|mandatory)\b", sentence, flags=re.I)][:20]
    scope = [sentence[:500] for sentence in sentences(text) if re.search(r"\b(scope|works include|scope of work|description of works)\b", sentence, flags=re.I)][:6]
    evaluation = [sentence[:500] for sentence in sentences(text) if re.search(r"\b(evaluation|technical score|financial score|points?|criteria)\b", sentence, flags=re.I)][:10]
    form_hits = matching_snippets(text, ["form ", "annexure", "schedule of return", "returnable"], limit=12)
    copies = re.findall(r"\b(\d{1,3})\s+(?:hard\s+)?copies\b", text, flags=re.I)
    return {
        "kind": "tender_document_review", "recommendation_only": True,
        "tender_number": tender.tender_number,
        "known_submission_deadline": as_utc(tender.submission_deadline).isoformat(),
        "document_date_hits": extract_date_hits(text), "required_documents": required,
        "mandatory_forms": form_hits, "technical_requirements": technical,
        "scope_summary_evidence": scope, "evaluation_criteria_evidence": evaluation,
        "copies_requested": int(copies[0]) if copies else None,
        "bid_security_referenced": any(item["name"] == "Bid security" for item in required),
        "explanation": "Text is extracted from the selected document and matched against published rule patterns. Staff must verify every finding against the tender document.",
    }


def latest_tender_review(db: Session, company_id: int, tender_id: int, document_id: int | None = None) -> dict[str, Any] | None:
    statement = select(AlgorithmicAssistantAnalysis).where(AlgorithmicAssistantAnalysis.company_id == company_id, AlgorithmicAssistantAnalysis.module == "tenders", AlgorithmicAssistantAnalysis.analysis_type == "tender_document_review", AlgorithmicAssistantAnalysis.subject_id == str(tender_id))
    if document_id is not None:
        statement = statement.where(AlgorithmicAssistantAnalysis.document_id == document_id)
    row = db.scalar(statement.order_by(AlgorithmicAssistantAnalysis.id.desc()).limit(1))
    return row.result if row else None


@router.get("/status")
def status(principal: Principal = Depends(current_principal)) -> dict[str, Any]:
    return {"status": "operational", "engine": "deterministic_rules", "rule_version": RULE_VERSION, "model_dependency": False, "capabilities": ["purchase_request_review", "quotation_comparison", "purchase_order_draft", "supplier_knowledge", "procurement_policy_answers", "tender_compliance_review", "tender_summary", "checklist_generation", "tender_questions", "draft_templates", "deadline_reminders"]}


@router.post("/procurement/requisitions/{requisition_id:int}/review")
def procurement_request_review(requisition_id: int, db: Session = Depends(get_db), principal: Principal = Depends(current_principal)) -> dict[str, Any]:
    row = requisition_or_404(db, principal, requisition_id, "procurement.manage")
    lines = db.scalars(select(ProcurementRequisitionLine).where(ProcurementRequisitionLine.requisition_id == row.id).order_by(ProcurementRequisitionLine.id)).all()
    result = review_requisition(row, lines)
    analysis = save_analysis(db, principal, module="procurement", analysis_type="purchase_request_review", subject_type="procurement_requisition", subject_id=row.id, branch_id=row.branch_id, site_id=row.site_id, document_id=None, source_sha256=None, result=result)
    return {"analysis": analysis, "result": result}


@router.post("/procurement/requisitions/{requisition_id:int}/quotation-comparison")
def procurement_quotation_comparison(requisition_id: int, db: Session = Depends(get_db), principal: Principal = Depends(current_principal)) -> dict[str, Any]:
    row = requisition_or_404(db, principal, requisition_id, "procurement.manage")
    result = quote_comparison(db, row)
    analysis = save_analysis(db, principal, module="procurement", analysis_type="quotation_comparison", subject_type="procurement_requisition", subject_id=row.id, branch_id=row.branch_id, site_id=row.site_id, document_id=None, source_sha256=None, result=result)
    return {"analysis": analysis, "result": result}


@router.get("/procurement/requisitions/{requisition_id:int}/purchase-order-draft")
def procurement_purchase_order_draft(requisition_id: int, delivery_location_id: int = Query(..., ge=1), db: Session = Depends(get_db), principal: Principal = Depends(current_principal)) -> dict[str, Any]:
    row = requisition_or_404(db, principal, requisition_id, "procurement.manage")
    quote = db.scalar(select(SupplierQuotation).where(SupplierQuotation.requisition_id == row.id, SupplierQuotation.status == "selected").order_by(SupplierQuotation.id.desc()).limit(1))
    if not quote:
        raise HTTPException(status_code=409, detail="Select a controlled supplier quotation before preparing a purchase-order draft")
    supplier = db.get(Supplier, quote.supplier_id)
    lines = db.scalars(select(SupplierQuotationLine).where(SupplierQuotationLine.quotation_id == quote.id).order_by(SupplierQuotationLine.id)).all()
    result = {
        "kind": "purchase_order_draft", "recommendation_only": True, "requisition_id": row.id,
        "quotation_id": quote.id, "supplier_id": quote.supplier_id, "supplier": supplier.name if supplier else None,
        "delivery_location_id": delivery_location_id, "order_date": date.today().isoformat(),
        "expected_delivery_date": (date.today() + timedelta(days=quote.delivery_days or 0)).isoformat() if quote.delivery_days is not None else None,
        "terms": f"Supplier payment terms: {supplier.payment_terms_days} days." if supplier else None,
        "lines": [{"requisition_line_id": line.requisition_line_id, "stock_item_id": line.stock_item_id, "quantity": str(line.quantity), "unit_price": str(line.unit_price), "line_total": str(line.line_total)} for line in lines],
        "total_amount": str(quote.total_amount),
        "next_step": "Review the draft, then create the official draft purchase order in Procurement & Stores. The assistant never creates or approves an order.",
    }
    analysis = save_analysis(db, principal, module="procurement", analysis_type="purchase_order_draft", subject_type="procurement_requisition", subject_id=row.id, branch_id=row.branch_id, site_id=row.site_id, document_id=quote.document_id, source_sha256=None, result=result)
    return {"analysis": analysis, "result": result}


@router.get("/procurement/supplier-knowledge")
def supplier_knowledge(question: str = Query(min_length=3, max_length=500), db: Session = Depends(get_db), principal: Principal = Depends(current_principal)) -> dict[str, Any]:
    require_anywhere(principal, "procurement.view")
    query = question.casefold()
    terms = [term for term in re.findall(r"[a-z0-9]{3,}", query) if term not in {"which", "supplier", "supplied", "supply", "offers", "offer", "best", "lowest", "average", "price", "last", "year", "have", "with", "from", "the", "for"}]
    item_term = terms[-1] if terms else ""
    cid = principal.user.company_id
    if "warranty" in query:
        candidates = []
        for quote in db.scalars(select(SupplierQuotation).where(SupplierQuotation.company_id == cid).order_by(SupplierQuotation.id.desc()).limit(250)).all():
            if not quote.document_id:
                continue
            try:
                text, _, _ = latest_document_text(db, cid, quote.document_id)
                warranty = extract_terms(text)["warranty_days"]
            except HTTPException:
                warranty = None
            if warranty is not None:
                supplier = db.get(Supplier, quote.supplier_id)
                candidates.append({"supplier": supplier.name if supplier else "Unknown", "quotation_id": quote.id, "warranty_days": warranty})
        candidates.sort(key=lambda item: (-int(item["warranty_days"]), str(item["supplier"])))
        return {"engine": "deterministic_rules", "answer": f"{candidates[0]['supplier']} has the longest extracted warranty in the recorded quotation documents." if candidates else "No readable recorded quotation warranty was found.", "evidence": candidates[:20], "limitation": "Only readable uploaded quotation documents are compared."}
    if "lowest" in query or "average price" in query:
        rows = db.execute(select(Supplier.name, func.avg(SupplierQuotationLine.unit_price).label("average_price"), func.count(SupplierQuotationLine.id).label("line_count")).join(SupplierQuotation, SupplierQuotation.supplier_id == Supplier.id).join(SupplierQuotationLine, SupplierQuotationLine.quotation_id == SupplierQuotation.id).join(ProcurementRequisitionLine, ProcurementRequisitionLine.id == SupplierQuotationLine.requisition_line_id).where(SupplierQuotation.company_id == cid, func.lower(ProcurementRequisitionLine.description).like(f"%{item_term}%") if item_term else True).group_by(Supplier.id, Supplier.name).order_by("average_price")).all()
        evidence = [{"supplier": name, "average_unit_price": str(price), "line_count": count} for name, price, count in rows]
        return {"engine": "deterministic_rules", "answer": f"{evidence[0]['supplier']} has the lowest recorded average unit price" + (f" for '{item_term}'." if item_term else ".") if evidence else "No matching quotation-line price history was found.", "evidence": evidence[:20], "limitation": "Prices are compared only where the requisition-line description matches the question term."}
    start = date(date.today().year - 1, 1, 1) if "last year" in query else date.today() - timedelta(days=365)
    rows = db.execute(select(Supplier.name, PurchaseOrder.purchase_order_number, PurchaseOrder.order_date, PurchaseOrderLine.description, PurchaseOrderLine.line_total).join(PurchaseOrder, PurchaseOrder.supplier_id == Supplier.id).join(PurchaseOrderLine, PurchaseOrderLine.purchase_order_id == PurchaseOrder.id).where(PurchaseOrder.company_id == cid, PurchaseOrder.order_date >= start, func.lower(PurchaseOrderLine.description).like(f"%{item_term}%") if item_term else True).order_by(PurchaseOrder.order_date.desc()).limit(50)).all()
    evidence = [{"supplier": name, "purchase_order": po, "order_date": order_date.isoformat(), "description": description, "line_total": str(total)} for name, po, order_date, description, total in rows]
    return {"engine": "deterministic_rules", "answer": f"{evidence[0]['supplier']} is the latest matching recorded supplier" + (f" for '{item_term}'." if item_term else ".") if evidence else "No matching purchase-order history was found.", "evidence": evidence, "limitation": "Results are based on recorded purchase-order lines, not an external supplier directory."}


@router.get("/procurement/policy-answer")
def procurement_policy_answer(question: str = Query(min_length=3, max_length=500), db: Session = Depends(get_db), principal: Principal = Depends(current_principal)) -> dict[str, Any]:
    require_anywhere(principal, "procurement.view")
    cfg = procurement_policy(db, principal.user.company_id)
    query = question.casefold()
    if any(word in query for word in ("quote", "quotation", "how many")):
        answer = f"{cfg['minimum_quotes']} quotations are required above M {cfg['low_value_quote_waiver']}."
    elif any(word in query for word in ("threshold", "waiver", "value")):
        answer = f"The low-value quotation waiver is M {cfg['low_value_quote_waiver']}; above it the {cfg['minimum_quotes']}-quotation rule applies."
    elif any(word in query for word in ("approval", "approve")):
        answer = "Requisitions and purchase orders follow their configured maker/checker approval workflows. The assistant cannot approve or bypass them."
    else:
        answer = "The current policy covers quotation count, low-value waiver, purchase-order overrun tolerance and reorder alerts. Ask about one of those controls for a direct rule answer."
    return {"engine": "deterministic_rules", "answer": answer, "policy": serialise(cfg), "source": "Company procurement policy"}


@router.post("/tenders/{tender_id:int}/document-review")
def tender_review_document(tender_id: int, payload: SourceInput, db: Session = Depends(get_db), principal: Principal = Depends(current_principal)) -> dict[str, Any]:
    tender = tender_or_404(db, principal, tender_id, "tenders.manage")
    text, document_id, digest, label = source_text(db, principal, payload)
    result = tender_document_review(text, tender)
    result["source_label"] = label
    analysis = save_analysis(db, principal, module="tenders", analysis_type="tender_document_review", subject_type="tender", subject_id=tender.id, branch_id=tender.branch_id, site_id=tender.site_id, document_id=document_id, source_sha256=digest, result=result)
    return {"analysis": analysis, "result": result}


@router.post("/tenders/{tender_id:int}/checklist-generate")
def tender_generate_checklist(tender_id: int, payload: SourceInput, db: Session = Depends(get_db), principal: Principal = Depends(current_principal)) -> dict[str, Any]:
    tender = tender_or_404(db, principal, tender_id, "tenders.manage")
    text, document_id, digest, label = source_text(db, principal, payload)
    review = tender_document_review(text, tender)
    existing = {item.name.casefold() for item in db.scalars(select(TenderChecklistItem).where(TenderChecklistItem.tender_id == tender.id)).all()}
    created = []
    for item in review["required_documents"]:
        if item["name"].casefold() in existing:
            continue
        row = TenderChecklistItem(company_id=tender.company_id, tender_id=tender.id, category=item["category"], name=item["name"], required=True, status="missing", due_date=tender.submission_deadline, notes="Generated by deterministic tender-document rules; verify against source.", updated_by=principal.user.full_name)
        db.add(row); db.flush(); created.append(row_dict(row)); existing.add(item["name"].casefold())
    result = {"kind": "tender_checklist_generation", "recommendation_only": True, "source_label": label, "created": created, "already_present_or_not_detected": len(review["required_documents"]) - len(created), "review": review}
    analysis = save_analysis(db, principal, module="tenders", analysis_type="tender_checklist_generation", subject_type="tender", subject_id=tender.id, branch_id=tender.branch_id, site_id=tender.site_id, document_id=document_id, source_sha256=digest, result=result)
    return {"analysis": analysis, "result": result}


@router.get("/tenders/{tender_id:int}/question")
def tender_question(tender_id: int, question: str = Query(min_length=3, max_length=500), document_id: int | None = None, db: Session = Depends(get_db), principal: Principal = Depends(current_principal)) -> dict[str, Any]:
    tender = tender_or_404(db, principal, tender_id)
    review = latest_tender_review(db, tender.company_id, tender.id, document_id)
    if not review:
        raise HTTPException(status_code=409, detail="Run Tender document review first so questions have a controlled rule-based source")
    query = question.casefold()
    evidence: list[Any] = []
    if any(term in query for term in ("closing", "deadline", "date", "close")):
        answer = f"The registered submission deadline is {review['known_submission_deadline']}."
        evidence = review.get("document_date_hits", [])
    elif any(term in query for term in ("bid bond", "bid security", "security")):
        answer = "The document rules detected a bid-security requirement." if review.get("bid_security_referenced") else "The rule scan did not detect a bid-security term; verify the tender document manually."
        evidence = [item for item in review.get("required_documents", []) if item["name"] == "Bid security"]
    elif "cop" in query:
        count = review.get("copies_requested")
        answer = f"The document text requests {count} copy/copies." if count is not None else "The rule scan did not find a number of copies; verify the submission instructions."
        evidence = review.get("mandatory_forms", [])
    elif any(term in query for term in ("experience", "method", "technical", "require")):
        evidence = [item for item in review.get("technical_requirements", []) if any(term in item.casefold() for term in ("experience", "method", "technical", "must", "shall", "required"))]
        answer = evidence[0] if evidence else "No matching technical requirement was extracted; verify the tender document manually."
    else:
        evidence = review.get("scope_summary_evidence", [])
        answer = evidence[0] if evidence else "Ask about the closing date, bid security, copies, experience or technical requirements."
    return {"engine": "deterministic_rules", "answer": answer, "evidence": evidence[:10], "source_analysis_id": None, "limitation": "This is a rule-based extraction aid. Staff must verify the original tender document."}


@router.post("/tenders/{tender_id:int}/draft")
def tender_draft(tender_id: int, payload: DraftInput, db: Session = Depends(get_db), principal: Principal = Depends(current_principal)) -> dict[str, Any]:
    tender = tender_or_404(db, principal, tender_id, "tenders.manage")
    review = latest_tender_review(db, tender.company_id, tender.id, payload.document_id)
    scope = (review or {}).get("scope_summary_evidence", [])
    requirements = (review or {}).get("technical_requirements", [])
    heading = {"company_profile": "Company Profile", "technical_response": "Technical Response", "method_statement": "Method Statement", "cover_letter": "Tender Cover Letter"}[payload.kind]
    if payload.kind == "cover_letter":
        draft = f"Re: {tender.title}\nTender reference: {tender.external_reference or tender.tender_number}\n\nDear Evaluation Committee,\n\nNthane Brothers submits its response for the above tender. The submission is subject to the controlled checklist, commercial approvals and tender documents recorded in BuildTrack.\n\nYours faithfully,\n[Authorised signatory]"
    elif payload.kind == "company_profile":
        draft = f"{heading}\n\nCompany: Nthane Brothers\nTender: {tender.title}\nClient: {tender.client_name}\n\n[Insert verified company registration, tax compliance, licences, relevant experience, resources and contact details. Do not invent credentials.]"
    elif payload.kind == "method_statement":
        draft = f"{heading}\n\nTender: {tender.title}\n\n1. Scope understanding\n{scope[0] if scope else '[Verify and insert scope from the tender document.]'}\n\n2. Delivery approach\n[Insert project-specific sequence, resources, quality controls, health and safety controls and programme.]\n\n3. Mandatory requirements to verify\n" + "\n".join(f"- {item}" for item in requirements[:8] or ["Verify tender requirements before submission."])
    else:
        draft = f"{heading}\n\nTender: {tender.title}\nClient: {tender.client_name}\n\nScope evidence\n{scope[0] if scope else '[Verify and insert scope from the tender document.]'}\n\nCompliance response\n" + "\n".join(f"- [Respond to] {item}" for item in requirements[:10] or ["[Add verified technical response items.]"])
    result = {"kind": "tender_draft", "recommendation_only": True, "draft_type": payload.kind, "title": heading, "content": draft, "notice": "This template contains placeholders and recorded tender facts only. Review, edit and approve it through document control before submission."}
    analysis = save_analysis(db, principal, module="tenders", analysis_type="tender_draft", subject_type="tender", subject_id=tender.id, branch_id=tender.branch_id, site_id=tender.site_id, document_id=payload.document_id, source_sha256=None, result=result)
    return {"analysis": analysis, "result": result}


@router.get("/tenders/{tender_id:int}/deadline-reminders")
def tender_deadline_reminders(tender_id: int, db: Session = Depends(get_db), principal: Principal = Depends(current_principal)) -> dict[str, Any]:
    tender = tender_or_404(db, principal, tender_id)
    deadline = as_utc(tender.submission_deadline)
    now = datetime.now(timezone.utc)
    reminders = []
    for days in (30, 14, 7, 1):
        trigger = deadline - timedelta(days=days)
        reminders.append({"days_remaining": days, "trigger_at": trigger.isoformat(), "status": "overdue" if now > deadline else "due" if now >= trigger else "scheduled", "message": f"{days} day{'s' if days != 1 else ''} remaining before tender submission deadline."})
    return {"engine": "deterministic_rules", "tender_number": tender.tender_number, "submission_deadline": deadline.isoformat(), "reminders": reminders, "delivery": "Shown in the assistant workspace and available to any configured notification scheduler; no external message is sent without an approved notification integration."}


@router.get("/history")
def history(module: Literal["procurement", "tenders"] | None = None, limit: int = Query(default=100, ge=1, le=500), db: Session = Depends(get_db), principal: Principal = Depends(current_principal)) -> list[dict[str, Any]]:
    if module == "procurement":
        require_anywhere(principal, "procurement.view")
    elif module == "tenders":
        require_anywhere(principal, "tenders.view")
    elif not (principal.has_permission_anywhere("procurement.view") or principal.has_permission_anywhere("tenders.view")):
        raise HTTPException(status_code=403, detail="Procurement or tender view permission is required")
    statement = select(AlgorithmicAssistantAnalysis).where(AlgorithmicAssistantAnalysis.company_id == principal.user.company_id)
    if module:
        statement = statement.where(AlgorithmicAssistantAnalysis.module == module)
    rows = db.scalars(statement.order_by(AlgorithmicAssistantAnalysis.id.desc()).limit(limit)).all()
    return [row_dict(row) for row in rows if row.branch_id is None or (principal.can("procurement.view", branch_id=row.branch_id, site_id=row.site_id) or principal.can("tenders.view", branch_id=row.branch_id, site_id=row.site_id))]
