from __future__ import annotations

import io
import json
from datetime import datetime, timezone
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

from core.access_control import platform_admin
from database.models import User
from integrations.mpesa.sandbox_matrix import (
    MPESA_DOCUMENTED_NOT_RUNNABLE,
    MPESA_OFFICIAL_SANDBOX_MATRIX,
)

router = APIRouter(
    prefix="/admin/testing/mpesa-certification",
    tags=["M-Pesa Testing Reports"],
)

SANDBOX_REPORT_PRODUCTS = tuple(MPESA_OFFICIAL_SANDBOX_MATRIX.keys())
LIVE_REPORT_PRODUCTS = (
    "c2b",
    "b2c",
    "b2b",
    "authorization",
    "query_transaction_status",
    "reversal",
    "update_transaction",
    "direct_debit_create",
    "direct_debit_payment",
    "query_direct_debit_reference",
    "query_direct_debit_customer",
    "query_direct_debit_mandate",
    "query_direct_debit_balance",
    "direct_debit_cancel",
)

SECTION_TITLES = {
    "c2b": "C2B Scenarios",
    "b2c": "B2C Scenarios",
    "b2b": "B2B Scenarios",
    "query_transaction_status": "Query Transaction Status Scenarios",
    "reversal": "Reversal Scenarios",
    "update_transaction": "Update Transaction Status Scenarios",
    "direct_debit_create": "Direct Debit Create Scenarios",
    "direct_debit_payment": "Direct Debit Payment Scenarios",
    "query_direct_debit_reference": "Query Direct Debit - ThirdPartyReference",
    "query_direct_debit_customer": "Query Direct Debit - Customer Status",
    "query_direct_debit_mandate": "Query Direct Debit - Mandate Status",
    "query_direct_debit_balance": "Query Direct Debit - Balance",
    "direct_debit_cancel": "Direct Debit Cancel Scenarios",
}

LIVE_PRODUCT_LABELS = {
    "c2b": "C2B Collection",
    "b2c": "B2C Payout",
    "b2b": "B2B Transfer",
    "authorization": "Two-stage C2B Authorization",
    "query_transaction_status": "Query Transaction Status",
    "reversal": "Reversal",
    "update_transaction": "Update Transaction Status",
    "direct_debit_create": "Direct Debit Create",
    "direct_debit_payment": "Direct Debit Payment",
    "query_direct_debit_reference": "Query Direct Debit - ThirdPartyReference",
    "query_direct_debit_customer": "Query Direct Debit - Customer Status",
    "query_direct_debit_mandate": "Query Direct Debit - Mandate Status",
    "query_direct_debit_balance": "Query Direct Debit - Balance",
    "direct_debit_cancel": "Direct Debit Cancel",
}


class MpesaTestingReportResult(BaseModel):
    product: str = Field(min_length=1, max_length=80)
    product_label: str | None = None
    scenario: str = Field(min_length=1, max_length=100)
    trigger_field: str | None = None
    trigger_value: str | None = None
    passed: bool = False
    duration_ms: float | None = None
    executed_at: str | None = None
    checks: dict[str, Any] = Field(default_factory=dict)
    request_evidence: dict[str, Any] = Field(default_factory=dict)
    result: dict[str, Any] = Field(default_factory=dict)


class MpesaTestingReportRequest(BaseModel):
    environment: str = Field(default="sandbox", max_length=30)
    shortcode: str | None = Field(default=None, max_length=32)
    session_requisition_passed: bool | None = None
    title: str = Field(default="M-Pesa OpenAPI Testing Report", max_length=160)
    results: list[MpesaTestingReportResult] = Field(default_factory=list, max_length=400)


def _is_sandbox(payload: MpesaTestingReportRequest) -> bool:
    return payload.environment.strip().lower() == "sandbox"


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_text(
    cell,
    text: str,
    *,
    bold: bool = False,
    color: str | None = None,
    size: float = 9.0,
    font: str = "Aptos",
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = font
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _json_block(cell, payload: dict[str, Any] | None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    text = json.dumps(payload or {}, indent=2, ensure_ascii=False)
    if not payload:
        text = "Not run / no provider evidence captured."
    run = paragraph.add_run(text)
    run.font.name = "Courier New"
    run.font.size = Pt(7.5)


def _configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.58)
    section.right_margin = Inches(0.58)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9.5)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(22)
    styles["Title"].font.bold = True
    styles["Heading 1"].font.name = "Aptos Display"
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 1"].font.bold = True
    styles["Heading 2"].font.name = "Aptos Display"
    styles["Heading 2"].font.size = Pt(12)
    styles["Heading 2"].font.bold = True


def _scenario_label(product: str, scenario: str) -> str:
    if scenario == "success_not_reversed":
        return "Valid Transaction Query"
    if scenario == "success_reversed":
        return "Valid Query (Reversed Transaction)"
    if scenario == "success":
        if product == "reversal":
            return "Valid Reversal Transaction"
        if product == "b2b":
            return "Valid B2B Payment"
        return "Valid Payment"
    if product == "c2b" and scenario == "transaction_failed":
        return "Transaction Failed [Incorrect PIN] (Alternate Trigger)"
    labels = {
        "request_timeout": "Request Timeout",
        "ussd_push_timeout": "USSD Push Timeout",
        "incorrect_pin": "Transaction Failed [Incorrect PIN]",
        "transaction_failed": "Transaction Failed",
        "internal_error": "Internal Error",
        "internal_error_2": "Internal Error (Alternate Trigger)",
        "invalid_amount": "Invalid Amount Used",
        "insufficient_balance": "Insufficient Balance",
        "service_unavailable": "Service Not Available",
        "missing_reference": "Missing Reference",
        "not_owned": "This transaction do not belong to you",
        "internal_error_long": "Internal Error (Long Reference)",
        "transaction_failed_long": "Transaction Failed (Long Reference)",
        "service_unavailable_long": "Service Not Available (Long Reference)",
        "invalid_use_case": "Invalid Use Case",
        "mandate_not_found": "Mandate does not exist",
        "pending_approval": "Pending Approval",
        "active": "Active",
        "locked": "Locked",
        "inactive": "Inactive",
        "cancelled": "Cancelled",
        "expired": "Expired",
        "larger_than_500": "Balance larger than 500",
        "smaller_than_500": "Balance smaller than 500",
        "live_verification": "Live Verification",
    }
    return labels.get(scenario, scenario.replace("_", " ").title())


def _latest_results(payload: MpesaTestingReportRequest) -> dict[tuple[str, str], MpesaTestingReportResult]:
    latest: dict[tuple[str, str], MpesaTestingReportResult] = {}
    for item in payload.results:
        latest[(item.product, item.scenario)] = item
    return latest


def _latest_live_results(payload: MpesaTestingReportRequest) -> dict[str, MpesaTestingReportResult]:
    latest: dict[str, MpesaTestingReportResult] = {}
    for item in payload.results:
        latest[item.product] = item
    return latest


def _add_cover(doc: Document, payload: MpesaTestingReportRequest, generated_by: str) -> None:
    brand = doc.add_paragraph()
    brand.alignment = WD_ALIGN_PARAGRAPH.CENTER
    brand_run = brand.add_run("ITHUTE PAY BRIDGE")
    brand_run.bold = True
    brand_run.font.name = "Aptos Display"
    brand_run.font.size = Pt(11)
    brand_run.font.color.rgb = RGBColor(15, 111, 189)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(payload.title)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("Provider testing evidence generated from captured PayBridge results")
    subtitle_run.font.size = Pt(10)
    subtitle_run.font.color.rgb = RGBColor(90, 105, 125)

    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    metadata = [
        ("Provider", "Vodacom M-Pesa OpenAPI"),
        ("Environment", payload.environment.strip().title()),
        ("Shortcode", payload.shortcode or "Not supplied"),
    ]
    if _is_sandbox(payload):
        metadata.append(("Session ID Requisition", "Passed" if payload.session_requisition_passed else "Not run / not captured"))
    metadata.extend(
        [
            ("Generated by", generated_by),
            ("Generated at", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")),
        ]
    )
    for label, value in metadata:
        cells = table.add_row().cells
        _set_cell_text(cells[0], label, bold=True, color="27415F")
        _shade(cells[0], "EDF4FA")
        _set_cell_text(cells[1], value)

    doc.add_paragraph()
    note = doc.add_paragraph()
    note_run = note.add_run("Result interpretation: ")
    note_run.bold = True
    if _is_sandbox(payload):
        note.add_run(
            "Passed means the selected official sandbox fixture was satisfied by the normalized status or by the exact deterministic provider outcome. "
            "The provider's raw response remains authoritative evidence and is printed for every executed case. Unrun cases remain visible as a completion checklist."
        )
    else:
        note.add_run(
            "Live verification records the actual request and provider response. Money-changing tests require explicit operator confirmation and are never compared with sandbox trigger values."
        )


def _status(item: MpesaTestingReportResult | None) -> str:
    if item is None:
        return "Not Run"
    return "Passed" if item.passed else "Failed"


def _status_color(status: str) -> str:
    return "137A3A" if status == "Passed" else "B42318" if status == "Failed" else "6B7280"


def _provider_outcome(item: MpesaTestingReportResult | None) -> str:
    if item is None:
        return "—"
    return str(
        item.result.get("response_description")
        or item.checks.get("actual_provider_outcome")
        or item.result.get("provider_response", {}).get("output_ResponseDesc")
        or "—"
    )


def _add_sandbox_summary(doc: Document, latest: dict[tuple[str, str], MpesaTestingReportResult]) -> None:
    doc.add_heading("Testing Summary", level=1)
    for product in SANDBOX_REPORT_PRODUCTS:
        spec = MPESA_OFFICIAL_SANDBOX_MATRIX.get(product)
        if not spec:
            continue
        doc.add_heading(SECTION_TITLES.get(product, f"{spec.get('label', product)} Scenarios"), level=2)
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        for cell, text in zip(table.rows[0].cells, ("End Point / Scenario", "Trigger", "Result", "Provider Outcome")):
            _set_cell_text(cell, text, bold=True, color="FFFFFF", size=8)
            _shade(cell, "0F6FBD")
        for scenario, case in (spec.get("cases") or {}).items():
            item = latest.get((product, scenario))
            row = table.add_row().cells
            _set_cell_text(row[0], _scenario_label(product, scenario), size=8)
            _set_cell_text(row[1], str(case.get("value") or ""), font="Courier New", size=7.5)
            status = _status(item)
            _set_cell_text(row[2], status, bold=True, color=_status_color(status), size=8)
            _set_cell_text(row[3], _provider_outcome(item), size=7.5)
        doc.add_paragraph()


def _add_result_meta(doc: Document, item: MpesaTestingReportResult | None) -> None:
    status = _status(item)
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(f"Result: {status}")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(_status_color(status))
    if item and item.duration_ms is not None:
        paragraph.add_run(f"  |  Duration: {item.duration_ms:.2f} ms")
    if item and item.executed_at:
        paragraph.add_run(f"  |  Executed: {item.executed_at}")
    if item and item.result.get("response_code"):
        paragraph.add_run(f"  |  Response code: {item.result.get('response_code')}")
    if item and item.result.get("response_description"):
        paragraph.add_run(f"  |  Provider: {item.result.get('response_description')}")


def _add_evidence_table(doc: Document, item: MpesaTestingReportResult | None) -> None:
    evidence = doc.add_table(rows=2, cols=2)
    evidence.style = "Table Grid"
    _set_cell_text(evidence.cell(0, 0), "Request", bold=True, color="FFFFFF")
    _shade(evidence.cell(0, 0), "082B4D")
    _set_cell_text(evidence.cell(0, 1), "Response", bold=True, color="FFFFFF")
    _shade(evidence.cell(0, 1), "082B4D")
    request_payload = item.request_evidence if item else {}
    response_payload = (item.result or {}).get("provider_response") if item else {}
    _json_block(evidence.cell(1, 0), request_payload)
    _json_block(evidence.cell(1, 1), response_payload or {})


def _add_sandbox_details(doc: Document, latest: dict[tuple[str, str], MpesaTestingReportResult]) -> None:
    doc.add_page_break()
    doc.add_heading("Detailed Request / Response Evidence", level=1)
    for product in SANDBOX_REPORT_PRODUCTS:
        spec = MPESA_OFFICIAL_SANDBOX_MATRIX.get(product)
        if not spec:
            continue
        doc.add_heading(SECTION_TITLES.get(product, f"{spec.get('label', product)} Scenarios"), level=1)
        for scenario, case in (spec.get("cases") or {}).items():
            item = latest.get((product, scenario))
            heading = doc.add_paragraph()
            heading_run = heading.add_run(_scenario_label(product, scenario))
            heading_run.bold = True
            heading_run.font.size = Pt(11)
            heading_run.font.color.rgb = RGBColor(8, 43, 77)
            trigger = heading.add_run(f"   [{case.get('value', '')}]")
            trigger.font.name = "Courier New"
            trigger.font.size = Pt(8)
            trigger.font.color.rgb = RGBColor(15, 111, 189)
            _add_result_meta(doc, item)
            if item and item.checks:
                checks = doc.add_paragraph()
                checks.add_run("Evaluation: ").bold = True
                checks.add_run(json.dumps(item.checks, ensure_ascii=False, sort_keys=True))
            _add_evidence_table(doc, item)
            doc.add_paragraph()


def _add_pending_provider_contracts(doc: Document) -> None:
    if not MPESA_DOCUMENTED_NOT_RUNNABLE:
        return
    doc.add_heading("Documented Provider Modules Awaiting Confirmed Wire Contract", level=1)
    paragraph = doc.add_paragraph(
        "These modules are listed in the supplied M-Pesa testing documentation but PayBridge does not manufacture a production request shape without a confirmed Lesotho OpenAPI contract. They remain explicitly reported rather than silently omitted."
    )
    paragraph.style = doc.styles["Normal"]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("Module", "Status", "Reason")):
        _set_cell_text(cell, text, bold=True, color="FFFFFF")
        _shade(cell, "6B7280")
    for product, detail in MPESA_DOCUMENTED_NOT_RUNNABLE.items():
        row = table.add_row().cells
        _set_cell_text(row[0], product.replace("_", " ").title())
        _set_cell_text(row[1], "Provider contract required", bold=True, color="9A6700")
        _set_cell_text(row[2], str(detail.get("reason") or "Not runnable"), size=8)


def _add_live_summary(doc: Document, payload: MpesaTestingReportRequest) -> None:
    doc.add_heading("Live Production Verification Coverage", level=1)
    latest = _latest_live_results(payload)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("Operation", "Reference / Input", "Result", "Response Code", "Duration")):
        _set_cell_text(cell, text, bold=True, color="FFFFFF", size=8)
        _shade(cell, "8B1E1E")
    for product in LIVE_REPORT_PRODUCTS:
        item = latest.get(product)
        row = table.add_row().cells
        _set_cell_text(row[0], LIVE_PRODUCT_LABELS[product], size=8)
        _set_cell_text(row[1], item.trigger_value if item and item.trigger_value else "—", font="Courier New", size=7.5)
        status = _status(item)
        _set_cell_text(row[2], status, bold=True, color=_status_color(status), size=8)
        _set_cell_text(row[3], str(item.result.get("response_code") if item else "—"), size=8)
        _set_cell_text(row[4], f"{item.duration_ms:.2f} ms" if item and item.duration_ms is not None else "—", size=8)


def _add_live_details(doc: Document, payload: MpesaTestingReportRequest) -> None:
    doc.add_page_break()
    doc.add_heading("Detailed Live Request / Response Evidence", level=1)
    latest = _latest_live_results(payload)
    for product in LIVE_REPORT_PRODUCTS:
        item = latest.get(product)
        title = LIVE_PRODUCT_LABELS[product]
        heading = doc.add_paragraph()
        run = heading.add_run(title)
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(139, 30, 30)
        if item and item.trigger_value:
            trigger = heading.add_run(f"   [{item.trigger_value}]")
            trigger.font.name = "Courier New"
            trigger.font.size = Pt(8)
        _add_result_meta(doc, item)
        _add_evidence_table(doc, item)
        doc.add_paragraph()


def build_mpesa_testing_report(payload: MpesaTestingReportRequest, *, generated_by: str) -> bytes:
    doc = Document()
    _configure_document(doc)
    _add_cover(doc, payload, generated_by)
    if _is_sandbox(payload):
        latest = _latest_results(payload)
        _add_sandbox_summary(doc, latest)
        _add_sandbox_details(doc, latest)
        _add_pending_provider_contracts(doc)
    else:
        _add_live_summary(doc, payload)
        _add_live_details(doc, payload)
        _add_pending_provider_contracts(doc)
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


@router.post("/report.docx")
def download_report(payload: MpesaTestingReportRequest, user: User = Depends(platform_admin)):
    report = build_mpesa_testing_report(payload, generated_by=user.full_name or user.email)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M")
    environment = payload.environment.strip().lower() or "sandbox"
    filename = f"M-Pesa-{environment}-testing-report-{stamp}.docx"
    return StreamingResponse(
        io.BytesIO(report),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
