from __future__ import annotations

import io

from docx import Document

from routers.mpesa_testing_report import (
    LIVE_REPORT_PRODUCTS,
    SANDBOX_REPORT_PRODUCTS,
    MpesaTestingReportRequest,
    MpesaTestingReportResult,
    build_mpesa_testing_report,
)


def _all_text(document: Document) -> str:
    chunks: list[str] = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                chunks.append(cell.text)
    return "\n".join(chunks)


def test_docx_report_contains_all_runnable_sandbox_modules_and_exact_evidence() -> None:
    payload = MpesaTestingReportRequest(
        environment="sandbox",
        shortcode="000000",
        session_requisition_passed=True,
        results=[
            MpesaTestingReportResult(
                product="b2b",
                product_label="B2B Single Stage",
                scenario="success",
                trigger_field="ReceiverPartyCode",
                trigger_value="000001",
                passed=True,
                duration_ms=1051.32,
                executed_at="2026-08-19T13:33:00+00:00",
                checks={
                    "status_matches": True,
                    "provider_outcome_matches": True,
                    "actual_provider_outcome": "Request processed successfully",
                },
                request_evidence={
                    "input_Amount": "25.00",
                    "input_ReceiverPartyCode": "000001",
                    "input_PrimaryPartyCode": "000000",
                    "input_Country": "LES",
                    "input_Currency": "LSL",
                    "input_ThirdPartyConversationID": "CERT-EXAMPLE",
                },
                result={
                    "status": "succeeded",
                    "response_code": "INS-0",
                    "response_description": "Request processed successfully",
                    "provider_response": {
                        "output_ResponseCode": "INS-0",
                        "output_ResponseDesc": "Request processed successfully",
                    },
                },
            )
        ],
    )

    binary = build_mpesa_testing_report(payload, generated_by="Platform Admin")
    assert binary.startswith(b"PK")
    document = Document(io.BytesIO(binary))
    text = _all_text(document)

    assert "M-Pesa OpenAPI Testing Report" in text
    assert "Testing Summary" in text
    assert "C2B Scenarios" in text
    assert "B2C Scenarios" in text
    assert "B2B Scenarios" in text
    assert "Query Transaction Status Scenarios" in text
    assert "Reversal Scenarios" in text
    assert "Update Transaction Status Scenarios" in text
    assert "Direct Debit Create Scenarios" in text
    assert "Direct Debit Payment Scenarios" in text
    assert "Query Direct Debit - ThirdPartyReference" in text
    assert "Query Direct Debit - Customer Status" in text
    assert "Query Direct Debit - Mandate Status" in text
    assert "Query Direct Debit - Balance" in text
    assert "Direct Debit Cancel Scenarios" in text
    assert len(SANDBOX_REPORT_PRODUCTS) >= 13
    assert "Valid B2B Payment" in text
    assert "000001" in text
    assert "Passed" in text
    assert "Not Run" in text
    assert '"input_ReceiverPartyCode": "000001"' in text
    assert '"output_ResponseCode": "INS-0"' in text
    assert "Detailed Request / Response Evidence" in text
    assert "Provider Outcome" in text
    assert "Evaluation:" in text
    assert "Documented Provider Modules Awaiting Confirmed Wire Contract" in text
    assert "Query Transaction Status V31" in text
    assert "Query Beneficiary Name" in text


def test_report_keeps_latest_result_for_duplicate_product_scenario() -> None:
    payload = MpesaTestingReportRequest(
        results=[
            MpesaTestingReportResult(product="reversal", scenario="success", passed=False),
            MpesaTestingReportResult(product="reversal", scenario="success", passed=True),
        ]
    )
    binary = build_mpesa_testing_report(payload, generated_by="Platform Admin")
    document = Document(io.BytesIO(binary))
    text = _all_text(document)
    assert "Valid Reversal Transaction" in text
    assert "Result: Passed" in text


def test_live_report_lists_every_supported_production_module_even_if_not_run() -> None:
    payload = MpesaTestingReportRequest(
        environment="production",
        shortcode="123456",
        title="M-Pesa OpenAPI Live Production Verification Report",
        results=[
            MpesaTestingReportResult(
                product="b2c",
                scenario="live_verification",
                trigger_field="CustomerMSISDN",
                trigger_value="26658001234",
                passed=True,
                duration_ms=840.25,
                request_evidence={
                    "input_Amount": "1.00",
                    "input_CustomerMSISDN": "26658001234",
                    "input_ServiceProviderCode": "123456",
                },
                result={
                    "status": "processing",
                    "response_code": "INS-0",
                    "response_description": "Request processed successfully",
                    "provider_response": {
                        "output_ResponseCode": "INS-0",
                        "output_ResponseDesc": "Request processed successfully",
                    },
                },
            )
        ],
    )

    binary = build_mpesa_testing_report(payload, generated_by="Platform Admin")
    document = Document(io.BytesIO(binary))
    text = _all_text(document)

    assert "M-Pesa OpenAPI Live Production Verification Report" in text
    assert "Live Production Verification Coverage" in text
    assert "B2C Payout" in text
    assert "26658001234" in text
    assert "840.25 ms" in text
    assert "Detailed Live Request / Response Evidence" in text
    assert '"input_ServiceProviderCode": "123456"' in text
    assert '"output_ResponseCode": "INS-0"' in text
    for label in (
        "C2B Collection",
        "B2C Payout",
        "B2B Transfer",
        "Two-stage C2B Authorization",
        "Query Transaction Status",
        "Reversal",
        "Update Transaction Status",
        "Direct Debit Create",
        "Direct Debit Payment",
        "Query Direct Debit - ThirdPartyReference",
        "Query Direct Debit - Customer Status",
        "Query Direct Debit - Mandate Status",
        "Query Direct Debit - Balance",
        "Direct Debit Cancel",
    ):
        assert label in text
    assert len(LIVE_REPORT_PRODUCTS) == 14
    assert "Not Run" in text
    assert "C2B Scenarios" not in text
    assert "Documented Provider Modules Awaiting Confirmed Wire Contract" in text
